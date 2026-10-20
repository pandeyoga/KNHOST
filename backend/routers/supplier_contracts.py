"""FASE D/E — Router `supplier_contracts` (kontrak mitra makloon & supplier).

RBAC (tanpa modul izin baru yang berlebihan):
  * lihat   → `supplier_contract:view`   (admin, manager, warehouse, purchasing-role)
  * kelola  → `supplier_contract:create|update|delete` (admin, manager)
  * kebijakan makloon (toleransi/susut default) → admin/manager (`require_role`)

Semua endpoint ber-prefix `/api` (aturan ingress) dan mengembalikan ARRAY/OBJEK telanjang
(kontrak respons KN) kecuali saat paginasi diminta (`?page=`).
"""
from typing import Any, Dict, List, Optional

from fastapi import APIRouter, HTTPException, Request

from db import db
from dependencies import audit, require_permission, require_role
from entity_scope import entity_ctx, resolve_list_scope
from pagination import envelope, get_page_params, is_paged
from schemas_contracts import (ContractStatusIn, MakloonPolicyIn, SupplierContractCreate,
                               SupplierContractPatch, TariffPreviewIn)
from services import contract_service as cs
from services import doc_refs_service as _refs  # INV-REF-04 — sapu tautan sebelum hapus
from services import makloon_claim_service as mcs
from services.uom_rules_service import UomRuleError

router = APIRouter(prefix="/api")


def _err(exc: Exception) -> HTTPException:
    return HTTPException(status_code=400, detail=str(exc))


# ── Daftar & statistik ───────────────────────────────────────────────
@router.get("/supplier-contracts")
async def list_contracts(request: Request, q: str = "", contract_type: str = "",
                         partner_id: str = "", process_type: str = "", product_id: str = "",
                         status: str = "", entity_id: Optional[str] = None,
                         limit: int = 200) -> Any:
    await require_permission(request, "supplier_contract", "view")
    ctx = await entity_ctx(request)
    flt: Dict[str, Any] = resolve_list_scope("supplier_contracts", {}, ctx, entity_id)
    for key, val in (("contract_type", contract_type), ("partner_id", partner_id),
                     ("process_type", process_type), ("product_id", product_id),
                     ("status", status)):
        if val:
            flt[key] = val
    if is_paged(request):
        page, size, term, sort = get_page_params(request)
        rows = await cs.list_contracts(flt, q=q or term, limit=size, skip=(page - 1) * size,
                                       sort=sort or "-created_at")
        total = await cs.count_contracts(flt) if not (q or term) else len(
            await cs.list_contracts(flt, q=q or term, limit=1000))
        return envelope(rows, total, page, size)
    return await cs.list_contracts(flt, q=q, limit=limit)


@router.get("/supplier-contracts/stats")
async def contract_stats(request: Request, contract_type: str = "",
                         entity_id: Optional[str] = None) -> Dict[str, Any]:
    await require_permission(request, "supplier_contract", "view")
    ctx = await entity_ctx(request)
    flt = resolve_list_scope("supplier_contracts", {}, ctx, entity_id)
    if contract_type:
        flt["contract_type"] = contract_type
    return await cs.stats(flt)


@router.get("/supplier-contracts/policy")
async def get_policy(request: Request) -> Dict[str, Any]:
    """FASE E-4 (E4.5) — kebijakan makloon badan usaha AKTIF (global + override-nya)."""
    await require_permission(request, "supplier_contract", "view")
    ctx = await entity_ctx(request)
    return await cs.get_settings(ctx.active_entity_id)


@router.put("/supplier-contracts/policy")
async def update_policy(payload: MakloonPolicyIn, request: Request) -> Dict[str, Any]:
    actor = await require_role(request, ["admin", "manager"])
    try:
        out = await cs.update_settings(payload.model_dump(exclude_none=True),
                                       actor=actor.get("name", ""))
    except cs.ContractError as exc:
        raise _err(exc) from exc
    await audit(actor.get("name", ""), "makloon_policy_updated", "settings", "makloon",
                payload.model_dump(exclude_none=True))
    return out


@router.post("/supplier-contracts/resolve")
async def resolve_contract(request: Request, partner_id: str = "", process_type: str = "",
                           product_id: str = "", input_product_id: str = "",
                           contract_type: str = "makloon") -> Dict[str, Any]:
    """Kontrak aktif paling spesifik untuk (mitra × proses × produk) — dipakai wizard."""
    await require_permission(request, "supplier_contract", "view")
    ctx = await entity_ctx(request)
    found = await cs.resolve_active(partner_id=partner_id, contract_type=contract_type,
                                    process_type=process_type, product_id=product_id,
                                    input_product_id=input_product_id,
                                    entity_id=ctx.active_entity_id)
    return {"contract": found, "found": bool(found)}


@router.post("/supplier-contracts/tariff-preview")
async def tariff_preview(payload: TariffPreviewIn, request: Request) -> Dict[str, Any]:
    """Simulasi tarif (angka antara terlihat) — tidak menyimpan apa pun."""
    await require_permission(request, "supplier_contract", "view")
    try:
        return await cs.tariff_preview(payload.model_dump())
    except (cs.ContractError, UomRuleError) as exc:
        raise _err(exc) from exc


@router.get("/supplier-contracts/{contract_id}")
async def get_contract(contract_id: str, request: Request) -> Dict[str, Any]:
    await require_permission(request, "supplier_contract", "view")
    doc = await cs.get_contract(contract_id)
    if not doc:
        raise HTTPException(status_code=404, detail="Kontrak tidak ditemukan")
    return doc


@router.post("/supplier-contracts")
async def create_contract(payload: SupplierContractCreate, request: Request) -> Dict[str, Any]:
    actor = await require_permission(request, "supplier_contract", "create")
    ctx = await entity_ctx(request)
    data = payload.model_dump()
    entity_id = data.get("entity_id") or ctx.active_entity_id
    try:
        doc = await cs.create_contract(data, entity_id=entity_id, actor=actor.get("name", ""))
    except cs.ContractError as exc:
        raise _err(exc) from exc
    await audit(actor.get("name", ""), "supplier_contract_created", "supplier_contract",
                doc["id"], {"contract_number": doc.get("contract_number"),
                            "partner_name": doc.get("partner_name"),
                            "tariff_basis": doc.get("tariff_basis")})
    return doc


@router.patch("/supplier-contracts/{contract_id}")
async def patch_contract(contract_id: str, payload: SupplierContractPatch,
                         request: Request) -> Dict[str, Any]:
    actor = await require_permission(request, "supplier_contract", "update")
    try:
        doc = await cs.patch_contract(contract_id, payload.model_dump(exclude_none=True),
                                      actor=actor.get("name", ""))
    except cs.ContractError as exc:
        raise _err(exc) from exc
    await audit(actor.get("name", ""), "supplier_contract_updated", "supplier_contract",
                contract_id, payload.model_dump(exclude_none=True))
    return doc


@router.post("/supplier-contracts/{contract_id}/status")
async def set_contract_status(contract_id: str, payload: ContractStatusIn,
                              request: Request) -> Dict[str, Any]:
    actor = await require_permission(request, "supplier_contract", "update")
    try:
        doc = await cs.set_status(contract_id, payload.status, reason=payload.reason,
                                  actor=actor.get("name", ""))
    except cs.ContractError as exc:
        raise _err(exc) from exc
    await audit(actor.get("name", ""), "supplier_contract_status", "supplier_contract",
                contract_id, {"status": payload.status, "reason": payload.reason})
    return doc


@router.delete("/supplier-contracts/{contract_id}")
async def delete_contract(contract_id: str, request: Request) -> Dict[str, Any]:
    actor = await require_permission(request, "supplier_contract", "delete")
    doc = await cs.get_contract(contract_id)
    if not doc:
        raise HTTPException(status_code=404, detail="Kontrak tidak ditemukan")
    used = await db.makloon_orders.count_documents({"steps.contract_id": contract_id})
    # FASE E — kontrak PEMBELIAN dipakai oleh baris Purchase Order (jejak harga & audit).
    used_po = await db.purchase_orders.count_documents(
        {"$or": [{"items.contract_id": contract_id}, {"contract_ids": contract_id}]})
    if used or used_po:
        where = " · ".join(x for x in [f"{used} order makloon" if used else "",
                                       f"{used_po} purchase order" if used_po else ""] if x)
        raise HTTPException(status_code=409,
                            detail=f"Kontrak dipakai {where} — nonaktifkan "
                                   "(status 'terminated') alih-alih menghapus.")
    # INV-REF-04 — order makloon menaut kontrak dua arah (`fulfills`). Pagar di atas
    # hanya memeriksa `steps.contract_id`, sedangkan tautannya lahir dari
    # `contract_id` tingkat atas — jadi sapu tautan balik sebelum menghapus, supaya
    # tidak ada dokumen hidup yang menunjuk kontrak hantu.
    await _refs.safe_unlink_all("supplier_contract", contract_id)
    await db[cs.COLL].delete_one({"id": contract_id})
    await audit(actor.get("name", ""), "supplier_contract_deleted", "supplier_contract",
                contract_id, {"contract_number": doc.get("contract_number")})
    return {"deleted": True, "id": contract_id}


# ── Skor mitra (turunan klaim · PS-11) ────────────────────────────────
@router.get("/makloon-partners/scorecard")
async def partner_scorecard(request: Request, entity_id: Optional[str] = None) -> List[Dict[str, Any]]:
    await require_permission(request, "makloon", "view")
    ctx = await entity_ctx(request)
    flt = resolve_list_scope("makloon_orders", {}, ctx, entity_id)
    return await mcs.partner_scorecard(flt)


async def _claim_recap(request: Request, entity_id: Optional[str], date_from: str, date_to: str, makloon_id: str,
                       rows_offset: int = 0, rows_limit: Optional[int] = 1000) -> Dict[str, Any]:
    from dependencies import require_any_permission
    from services import makloon_recap_service as mrs
    await require_any_permission(request, [("makloon_order", "record_loss"), ("makloon_order", "approve_loss"),
                                           ("makloon_order", "claim_approve"), ("gl", "view"),
                                           ("vendor_bill", "view"), ("accounting", "view")])
    ctx = await entity_ctx(request)
    flt = resolve_list_scope("makloon_orders", {}, ctx, entity_id)
    return await mrs.recap(flt, date_from=date_from[:10], date_to=date_to[:10], makloon_id=makloon_id,
                           rows_offset=rows_offset, rows_limit=rows_limit)


@router.get("/makloon-claims/recap")
async def makloon_claim_recap(request: Request, entity_id: Optional[str] = None, date_from: str = "",
                              date_to: str = "", makloon_id: str = "", rows_offset: int = 0,
                              rows_limit: int = 1000) -> Dict[str, Any]:
    """L01 — rekap susut, klaim & potongan tagihan per mitra makloon (penagihan & evaluasi).
    AUD-MKL-02 — rincian per halaman (`rows_offset`, `rows_limit` maks 1.000) + `rows_total`; ringkasan selalu utuh."""
    if rows_offset < 0 or not 1 <= rows_limit <= 1000:
        raise HTTPException(status_code=400, detail="rows_offset minimal 0 dan rows_limit 1-1.000.")
    return await _claim_recap(request, entity_id, date_from, date_to, makloon_id, rows_offset, rows_limit)


ROWS_PER_SHEET = 100_000   # AUD-MKL-02 — rincian lengkap dibagi ke beberapa sheet bila sangat besar


@router.get("/makloon-claims/recap.xlsx")
async def makloon_claim_recap_xlsx(request: Request, entity_id: Optional[str] = None, date_from: str = "",
                                   date_to: str = "", makloon_id: str = ""):
    """L01 — unduh Excel rekap susut & klaim per mitra (lampiran penagihan bulanan).
    AUD-MKL-01 — kolom memakai kontrak data JSON rekap (`loss_claim`, `loss_deducted`); AUD-MKL-02 — SEMUA rincian."""
    import io
    from fastapi.responses import Response
    from openpyxl import Workbook
    from openpyxl.styles import Font
    data = await _claim_recap(request, entity_id, date_from, date_to, makloon_id, 0, None)
    wb = Workbook()
    ws = wb.active
    ws.title = "Ringkasan per Mitra"
    period = f"{date_from or 'awal'} s/d {date_to or 'sekarang'}"
    ws.append([f"Rekap Susut & Klaim Mitra Makloon — periode {period}"])
    ws["A1"].font = Font(bold=True, size=12)
    head = ["Mitra", "Kejadian susut", "Qty susut", "Beban susut wajar (Rp)", "Piutang klaim (Rp)",
            "Dipotong dari tagihan (Rp)", "Klaim selisih (jml)", "Klaim selisih (Rp)", "Susut menunggu",
            "Klaim selisih menunggu", "Total ditagih ke mitra (Rp)"]
    ws.append(head)
    for p in data["partners"]:
        ws.append([p["makloon_name"], p["loss_events"], p["loss_qty"], p["loss_expense"], p["loss_claim"],
                   p["loss_deducted"], p["variance_claims"], p["variance_claim_amount"], p["loss_pending"],
                   p["variance_pending"], round(p["receivable_total"] + p["deducted_total"], 2)])
    t = data["totals"]
    ws.append(["TOTAL", t["loss_events"], t["loss_qty"], t["loss_expense"], t["loss_claim"], t["loss_deducted"],
               t["variance_claims"], t["variance_claim_amount"], t["loss_pending"], t["variance_pending"],
               round(t["receivable_total"] + t["deducted_total"], 2)])
    for c in ws[2] + ws[ws.max_row]:
        c.font = Font(bold=True)
    ws.append([])
    ws.append([f"Rincian: {data['rows_total']} baris (lengkap) di sheet Rincian."])
    kind = {"susut": "Susut/hilang", "selisih": "Klaim selisih"}
    from db import db as _db
    bids = list({r["bill_id"] for r in data["rows"] if r.get("bill_id")})
    bno = {b["id"]: b.get("bill_number") or b["id"] async for b in _db.vendor_bills.find(
        {"id": {"$in": bids}}, {"_id": 0, "id": 1, "bill_number": 1})}
    settle = {"tagih_ganti": "Tagih ganti rugi", "potong_bon": "Potong tagihan jasa"}
    rows = data["rows"]
    sheets = []
    for i in range(0, max(len(rows), 1), ROWS_PER_SHEET):
        ws2 = wb.create_sheet("Rincian" if i == 0 else f"Rincian ({i // ROWS_PER_SHEET + 1})")
        ws2.append(["Tanggal", "Jenis", "No. MKO", "Langkah", "Mitra", "Qty", "Satuan", "Beban (Rp)", "Klaim (Rp)",
                    "Penyelesaian", "No. tagihan dipotong", "Status", "Alasan"])
        for c in ws2[1]:
            c.font = Font(bold=True)
        for r in rows[i:i + ROWS_PER_SHEET]:
            ws2.append([(r.get("at") or "")[:10], kind.get(r.get("kind"), r.get("kind")), r.get("mko_number"),
                        r.get("step_seq"), r.get("makloon_name"), r.get("qty"), r.get("unit"), r.get("expense"),
                        r.get("claim"), settle.get(r.get("settle"), r.get("settle") or ""), bno.get(r.get("bill_id"), ""),
                        r.get("status"), r.get("reason")])
        sheets.append(ws2)
    for sheet in [ws] + sheets:
        for col in sheet.iter_cols(max_row=min(sheet.max_row, 500)):
            width = max(len(str(c.value or "")) for c in col[1:] or col)
            sheet.column_dimensions[col[0].column_letter].width = min(max(10, width + 2), 50)
    buf = io.BytesIO()
    wb.save(buf)
    name = f"rekap_klaim_mitra_{(date_from or 'awal')}_{(date_to or 'sekarang')}.xlsx"
    return Response(content=buf.getvalue(),
                    media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f"attachment; filename={name}"})
