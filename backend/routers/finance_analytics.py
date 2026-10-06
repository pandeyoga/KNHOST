"""FINANCE — Analitik lanjutan: Profitabilitas, Proyeksi Kas, Control Tower.

Akses: permission module "accounting" (admin/manager). Respons OBJEK telanjang
(kontrak KN3). Semua ter-scope per entitas (buku terpisah per PT, F0-E).

Endpoint:
- GET /api/finance/profitability      → margin per produk/kategori/pelanggan/sales (WAC)
- GET /api/finance/cashflow-forecast  → proyeksi likuiditas (AR/AP jatuh tempo)
- GET /api/finance/tower              → dashboard keuangan terpadu (Control Tower)
"""
from typing import Any, Dict, Optional

from fastapi import APIRouter, Request, Query

from dependencies import require_permission
from core_utils import strip_cost_fields
from entity_scope import entity_ctx, resolve_list_scope
from services import profitability_service as prof
from services import cashflow_forecast_service as fc
from services import finance_tower_service as tower

router = APIRouter(prefix="/api")


async def _scope(request: Request, collection: str, entity_id: Optional[str]) -> Dict[str, Any]:
    ctx = await entity_ctx(request)
    return resolve_list_scope(collection, {}, ctx, entity_id)


@router.get("/finance/profitability")
async def get_profitability(
    request: Request,
    start: Optional[str] = Query(None),
    end: Optional[str] = Query(None),
    entity_id: Optional[str] = Query(None),
) -> Dict[str, Any]:
    """Analisis profitabilitas/margin (WAC) per 4 dimensi + tren bulanan."""
    actor = await require_permission(request, "accounting", "view")
    scope = await _scope(request, "sales_orders", entity_id)
    ent = entity_id if entity_id and entity_id != "all" else None
    # Tanya KN F0.4 — HPP/margin hanya untuk admin/manager (finance menerima pendapatan & qty).
    return strip_cost_fields(await prof.profitability(start=start, end=end, scope=scope, entity_id=ent),
                             actor.get("role"))


@router.get("/finance/profitability/export.xlsx")
async def export_profitability(
    request: Request,
    start: Optional[str] = Query(None),
    end: Optional[str] = Query(None),
    entity_id: Optional[str] = Query(None),
):
    """Excel: Realisasi (terkirim) vs Estimasi (nilai pesanan) berdampingan per pelanggan & produk."""
    import io
    from fastapi.responses import Response
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from core_utils import COST_SAFE_ROLES
    actor = await require_permission(request, "accounting", "view")
    scope = await _scope(request, "sales_orders", entity_id)
    ent = entity_id if entity_id and entity_id != "all" else None
    data = await prof.profitability(start=start, end=end, scope=scope, entity_id=ent)
    cost_ok = actor.get("role") in COST_SAFE_ROLES

    def _cell(v: Any) -> Any:
        return "'" + v if isinstance(v, str) and v[:1] in ("=", "+", "-", "@", "\t", "\r") else v

    def _sheet(ws, dim: str, label: str) -> None:
        real = {r["key"]: r for r in data.get(dim, [])}
        est = {r["key"]: r for r in (data.get("estimate") or {}).get(dim, [])}
        head = [label, "Realisasi Pendapatan"] + (["Realisasi HPP", "Realisasi Marjin", "Realisasi Marjin %"] if cost_ok else []) \
            + ["Estimasi Pendapatan"] + (["Estimasi HPP", "Estimasi Marjin"] if cost_ok else []) \
            + ["Selisih Pendapatan (Realisasi − Estimasi)", "Pesanan Realisasi", "Pesanan Estimasi"]
        ws.append(head)
        for c in ws[1]:
            c.font = Font(bold=True)
        keys = sorted(set(real) | set(est), key=lambda k: -float((real.get(k) or est.get(k) or {}).get("revenue") or 0))
        for k in keys:
            r, e = real.get(k, {}), est.get(k, {})
            row = [_cell(r.get("name") or e.get("name") or k), r.get("revenue", 0.0)]
            if cost_ok:
                row += [r.get("cogs", 0.0), r.get("margin", 0.0), r.get("margin_pct")]
            row.append(e.get("revenue", 0.0))
            if cost_ok:
                row += [e.get("cogs", 0.0), e.get("margin", 0.0)]
            row += [round(float(r.get("revenue", 0) or 0) - float(e.get("revenue", 0) or 0), 2),
                    r.get("orders", 0), e.get("orders", 0)]
            ws.append(row)
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = 22

    wb = Workbook()
    ws = wb.active
    ws.title = "Per Pelanggan"
    _sheet(ws, "by_customer", "Pelanggan")
    _sheet(wb.create_sheet("Per Produk"), "by_product", "Produk")
    info = wb.create_sheet("Keterangan")
    for row in (["Periode", f"{start or '-'} s/d {end or '-'}"],
                ["Realisasi", data.get("metric_label", "")],
                ["Estimasi", (data.get("estimate") or {}).get("label", "")],
                ["Pesanan lama tanpa surat jalan", data.get("legacy_orders", 0)],
                ["HPP/Marjin", "ditampilkan" if cost_ok else "disembunyikan untuk peran Anda"]):
        info.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    name = f"profitabilitas_{start or 'awal'}_{end or 'akhir'}.xlsx"
    return Response(content=buf.getvalue(),
                    media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="{name}"'})


@router.get("/finance/cashflow-forecast")
async def get_cashflow_forecast(
    request: Request,
    entity_id: Optional[str] = Query(None),
) -> Dict[str, Any]:
    """Proyeksi arus kas ke depan dari piutang & hutang jatuh tempo."""
    await require_permission(request, "accounting", "view")
    scope = await _scope(request, "sales_orders", entity_id)
    ent = entity_id if entity_id and entity_id != "all" else None
    return await fc.cashflow_forecast(scope=scope, entity_id=ent)


@router.get("/finance/tower")
async def get_finance_tower(
    request: Request,
    entity_id: Optional[str] = Query(None),
) -> Dict[str, Any]:
    """Dashboard Keuangan terpadu (Control Tower)."""
    await require_permission(request, "accounting", "view")
    ctx = await entity_ctx(request)
    scope = resolve_list_scope("journal_entries", {}, ctx, entity_id)
    # G3 D4-FIN-01 — tanpa parameter = entitas AKTIF (sama dengan jurnal), bukan gabungan semua entitas
    if entity_id == "all" or (not entity_id and getattr(ctx, "view_all", False)):
        ent, comp_ids = None, list(ctx.allowed_entity_ids)
    else:
        ent = entity_id or ctx.active_entity_id
        comp_ids = [ent]
    return await tower.finance_tower(scope=scope, entity_id=ent, entity_ids=comp_ids)
