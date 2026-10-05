"""Tanya KN F1 — endpoint lapisan semantik (dipakai tool server F2 & galeri template F3/F4)."""
from typing import Any, Dict, List, Optional

from fastapi import APIRouter, HTTPException, Query, Request

from db import db
from dependencies import audit, current_user, require_role
from entity_scope import entity_ctx
from services import analytics_catalog as cat
from services import analytics_engine as engine
from services import analytics_facts as facts

router = APIRouter(prefix="/api/ai")

AI_ROLES = ("admin", "manager", "sales_admin", "finance", "md", "warehouse_admin", "sales")


async def _ai_user(request: Request) -> Dict[str, Any]:
    from services.config_resolver import value_of
    user = await current_user(request)
    roles = await value_of("ai.enabled_roles") or list(AI_ROLES)
    if user.get("role") != "admin" and user.get("role") not in roles:
        raise HTTPException(status_code=403, detail="Asisten analitik belum dibuka untuk peran Anda")
    return user


async def _run(fn, payload: Dict[str, Any], request: Request) -> Dict[str, Any]:
    user = await _ai_user(request)
    try:
        return await fn(payload, user, await entity_ctx(request), session_id=payload.get("session_id"))
    except engine.AnalyticsError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc


@router.get("/catalog")
async def get_catalog(request: Request) -> Dict[str, Any]:
    await _ai_user(request)
    return {"version": cat.CATALOG_VERSION, "dimensions": cat.DIMENSIONS,
            "metrics": {k: {"label": v["label"], "unit": v["unit"], "definition": v["definition"],
                            "date_basis": v["date_basis"], "dimensions": sorted(cat.supported_dims(k))}
                        for k, v in cat.METRICS.items()}}


@router.post("/query")
async def post_query(payload: Dict[str, Any], request: Request) -> Dict[str, Any]:
    return await _run(engine.query_metrics, payload, request)


@router.post("/analyze-change")
async def post_analyze_change(payload: Dict[str, Any], request: Request) -> Dict[str, Any]:
    return await _run(engine.analyze_change, payload, request)


@router.get("/results/{result_id}")
async def get_result(result_id: str, request: Request) -> Dict[str, Any]:
    user = await _ai_user(request)
    doc = await db.ai_results.find_one({"result_id": result_id}, {"_id": 0, "expires_at": 0})
    if not doc or (doc.get("user_id") != user["id"] and user.get("role") != "admin"):
        raise HTTPException(status_code=404, detail="Hasil tidak ditemukan atau sudah kedaluwarsa")
    return doc


@router.get("/facts/status")
async def facts_status(request: Request) -> Dict[str, Any]:
    await require_role(request, ["admin", "manager"])
    st = await db.ai_fact_state.find_one({"id": facts.STATE_ID}, {"_id": 0}) or {}
    return {**st, "rows": await db.fact_sales_lines.count_documents({})}


@router.post("/facts/rebuild")
async def facts_rebuild(request: Request, date_from: Optional[str] = Query(None, alias="from")) -> Dict[str, Any]:
    actor = await require_role(request, ["admin"])
    res = await facts.rebuild_facts(date_from)
    await audit(actor.get("name", ""), "ai_facts_rebuild", "fact_sales_lines", facts.STATE_ID, res)
    return res


# ── F2/F3 — template, chat, sesi, unduh ──────────────────────────────────────
@router.get("/status")
async def get_status(request: Request) -> Dict[str, Any]:
    from services import ai_chat_service
    return await ai_chat_service.status(await _ai_user(request))


@router.get("/templates")
async def list_templates(request: Request) -> Dict[str, Any]:
    from services import ai_templates
    user = await _ai_user(request)
    return {"version": ai_templates.TEMPLATES["version"], "groups": ai_templates.GROUP_ORDER,
            "period_presets": ai_templates.PERIOD_PRESETS, "templates": ai_templates.for_role(user.get("role") or "")}


@router.post("/templates/{template_id}/run")
async def run_template(template_id: str, payload: Dict[str, Any], request: Request) -> Dict[str, Any]:
    from services import ai_templates
    user = await _ai_user(request)
    try:
        return await ai_templates.run(template_id, payload.get("params") or {}, user, await entity_ctx(request))
    except engine.AnalyticsError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc


@router.post("/tools/{tool}")
async def run_tool(tool: str, payload: Dict[str, Any], request: Request) -> Dict[str, Any]:
    from services import ai_tools
    user = await _ai_user(request)
    res = await ai_tools.execute(tool, payload, user, await entity_ctx(request))
    if res.get("error"):
        raise HTTPException(status_code=400, detail=res["error"])
    return res


@router.post("/chat")
async def post_chat(payload: Dict[str, Any], request: Request):
    import json as _json
    from fastapi.responses import StreamingResponse
    from services import ai_chat_service
    user = await _ai_user(request)
    ctx = await entity_ctx(request)
    question = str(payload.get("question") or "").strip()
    if not question:
        raise HTTPException(status_code=400, detail="Pertanyaan kosong")

    async def gen():
        try:
            async for ev in ai_chat_service.chat_events(question, payload.get("session_id"), bool(payload.get("deep")), user, ctx,
                                                        payload.get("context")):
                yield f"data: {_json.dumps(ev, ensure_ascii=False, default=str)}\n\n"
        except Exception as exc:  # noqa: BLE001 — galat model dikirim sebagai event, bukan koneksi putus
            yield f"data: {_json.dumps({'type': 'done', 'error': f'Asisten gagal menjawab: {exc}'})}\n\n"
    return StreamingResponse(gen(), media_type="text/event-stream",
                             headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no"})


@router.get("/sessions")
async def list_sessions(request: Request, q: str = "", limit: int = 50) -> List[Dict[str, Any]]:
    import re
    user = await _ai_user(request)
    match: Dict[str, Any] = {"user_id": user["id"]}
    if q.strip():
        match["$or"] = [{"title": {"$regex": re.escape(q.strip()), "$options": "i"}},
                        {"turns.content": {"$regex": re.escape(q.strip()), "$options": "i"}}]
    return await db.ai_chat_sessions.aggregate([
        {"$match": match}, {"$sort": {"updated_at": -1}}, {"$limit": max(1, min(limit, 100))},
        {"$project": {"_id": 0, "id": 1, "title": 1, "updated_at": 1, "created_at": 1,
                      "questions": {"$size": {"$filter": {"input": {"$ifNull": ["$turns", []]}, "cond": {"$eq": ["$$this.role", "user"]}}}}}}]).to_list(100)


@router.delete("/sessions/{session_id}")
async def delete_session(session_id: str, request: Request) -> Dict[str, Any]:
    user = await _ai_user(request)
    res = await db.ai_chat_sessions.delete_one({"id": session_id, "user_id": user["id"]})
    if not res.deleted_count:
        raise HTTPException(status_code=404, detail="Sesi tidak ditemukan")
    return {"ok": True}


@router.get("/sessions/{session_id}")
async def get_session(session_id: str, request: Request) -> Dict[str, Any]:
    user = await _ai_user(request)
    q = {"id": session_id} if user.get("role") == "admin" else {"id": session_id, "user_id": user["id"]}
    doc = await db.ai_chat_sessions.find_one(q, {"_id": 0, "items": 0})
    if not doc:
        raise HTTPException(status_code=404, detail="Sesi tidak ditemukan")
    return doc


@router.post("/feedback")
async def post_feedback(payload: Dict[str, Any], request: Request) -> Dict[str, Any]:
    from core_utils import new_id, now_iso
    user = await _ai_user(request)
    if payload.get("rating") not in ("up", "down"):
        raise HTTPException(status_code=400, detail="Penilaian harus up atau down")
    doc = {"id": new_id("aifb"), "user_id": user["id"], "rating": payload["rating"],
           "comment": str(payload.get("comment") or "")[:1000], "session_id": payload.get("session_id") or "",
           "template_id": payload.get("template_id") or "", "result_ids": payload.get("result_ids") or [], "created_at": now_iso()}
    await db.ai_feedback.insert_one(dict(doc))
    return {"ok": True, "id": doc["id"]}


_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _fill_sheet(ws, doc: Dict[str, Any]) -> None:
    from openpyxl.styles import Font
    ws.append([_safe_cell(doc.get("title") or "Tanya KN")])
    ws["A1"].font = Font(bold=True, size=13)
    per = doc.get("period") or {}
    ws.append([f"Periode: {per.get('label', '-')}"])
    ws.append([f"Definisi: {doc.get('definition') or '-'}"])
    ws.append([])
    cols = doc.get("columns") or []
    ws.append([c["label"] for c in cols])
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)
    fmt = {"idr": "#,##0", "qty": "#,##0.00", "pct": "0.00", "count": "#,##0"}
    for r in doc.get("rows") or []:
        ws.append([_safe_cell(r.get(c["key"])) for c in cols])
        for i, c in enumerate(cols, start=1):
            if c.get("unit") in fmt:
                ws.cell(row=ws.max_row, column=i).number_format = fmt[c["unit"]]
    for i in range(1, len(cols) + 1):
        ws.column_dimensions[ws.cell(row=5, column=i).column_letter].width = 22


def _safe_cell(v: Any) -> Any:
    """KN-E32 — teks yang diawali = + - @ (atau tab/CR) diawali apostrof agar tidak jadi formula."""
    if isinstance(v, str) and v[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + v
    return v


def _xlsx_response(wb, name: str):
    import io
    from fastapi.responses import Response
    buf = io.BytesIO()
    wb.save(buf)
    return Response(content=buf.getvalue(), media_type=_XLSX,
                    headers={"Content-Disposition": f'attachment; filename="{name}.xlsx"'})


@router.get("/results/{result_id}/export.xlsx")
async def export_result(result_id: str, request: Request):
    from openpyxl import Workbook
    doc = await get_result(result_id, request)
    wb = Workbook()
    wb.active.title = "Hasil"
    _fill_sheet(wb.active, doc)
    return _xlsx_response(wb, f"tanya-kn-{result_id}")


@router.get("/export.xlsx")
async def export_answer(request: Request, ids: str = ""):
    """Satu jawaban (template/chat) → satu workbook, satu sheet per hasil."""
    import re
    from openpyxl import Workbook
    rids = list(dict.fromkeys(x for x in ids.split(",") if x.strip()))[:20]
    if not rids:
        raise HTTPException(status_code=400, detail="Tidak ada hasil untuk diunduh")
    wb = Workbook()
    wb.remove(wb.active)
    used: set = set()
    for n, rid in enumerate(rids, start=1):
        doc = await get_result(rid.strip(), request)
        name = re.sub(r"[\\/*?:\[\]]", " ", doc.get("title") or f"Hasil {n}")[:28].strip() or f"Hasil {n}"
        if name in used:
            name = f"{name[:25]} {n}"
        used.add(name)
        _fill_sheet(wb.create_sheet(name), doc)
    return _xlsx_response(wb, "tanya-kn-jawaban")


# ── F4 — narasi AI, template pribadi, laporan terjadwal · F5 — snapshot · biaya ─────
async def _owned_results(ids: List[str], user) -> List[Dict[str, Any]]:
    q = {"result_id": {"$in": ids[:20]}} if user.get("role") == "admin" else {"result_id": {"$in": ids[:20]}, "user_id": user["id"]}
    docs = {d["result_id"]: d async for d in db.ai_results.find(q, {"_id": 0, "expires_at": 0})}
    return [docs[i] for i in ids if i in docs]


@router.post("/narrative")
async def post_narrative(payload: Dict[str, Any], request: Request) -> Dict[str, Any]:
    from services import ai_templates
    from services.ai_narrative import narrate
    user = await _ai_user(request)
    tpl = ai_templates.BY_ID.get(payload.get("template_id") or "")
    if not tpl or (user.get("role") != "admin" and user.get("role") not in tpl["roles"]):
        raise HTTPException(status_code=404, detail="Template tidak ditemukan")
    results = await _owned_results([str(x) for x in payload.get("result_ids") or []], user)
    if not results:
        raise HTTPException(status_code=404, detail="Hasil tidak ditemukan atau sudah kedaluwarsa")
    # KN-E44 — narasi AI ikut kuota harian; kuota habis → ringkasan otomatis tanpa AI.
    from services.ai_chat_service import _questions_today, chat_config
    allow_ai = await _questions_today(user["id"]) < int((await chat_config())["daily_question_limit"] or 60)
    return await narrate(tpl, results, user, allow_ai=allow_ai) or {"text": None, "source": "none"}


@router.get("/my-templates")
async def list_my_templates(request: Request) -> List[Dict[str, Any]]:
    user = await _ai_user(request)
    return await db.ai_user_templates.find({"user_id": user["id"]}, {"_id": 0}).sort("created_at", -1).to_list(50)


@router.post("/my-templates")
async def create_my_template(payload: Dict[str, Any], request: Request) -> Dict[str, Any]:
    from core_utils import new_id, now_iso
    user = await _ai_user(request)
    question = str(payload.get("question") or "").strip()
    if len(question) < 3:
        raise HTTPException(status_code=400, detail="Pertanyaan template terlalu pendek")
    if await db.ai_user_templates.count_documents({"user_id": user["id"]}) >= 50:
        raise HTTPException(status_code=400, detail="Maksimal 50 template pribadi")
    doc = {"id": new_id("aiut"), "user_id": user["id"], "mode": "llm", "question": question[:500],
           "title": (str(payload.get("title") or "").strip() or question)[:80], "created_at": now_iso()}
    await db.ai_user_templates.insert_one(dict(doc))
    return doc


@router.delete("/my-templates/{template_id}")
async def delete_my_template(template_id: str, request: Request) -> Dict[str, Any]:
    user = await _ai_user(request)
    if not (await db.ai_user_templates.delete_one({"id": template_id, "user_id": user["id"]})).deleted_count:
        raise HTTPException(status_code=404, detail="Template tidak ditemukan")
    return {"ok": True}


async def _own_schedule(schedule_id: str, user) -> Dict[str, Any]:
    q = {"id": schedule_id} if user.get("role") == "admin" else {"id": schedule_id, "user_id": user["id"]}
    doc = await db.ai_schedules.find_one(q, {"_id": 0})
    if not doc:
        raise HTTPException(status_code=404, detail="Jadwal tidak ditemukan")
    return doc


@router.get("/schedules")
async def list_schedules(request: Request) -> List[Dict[str, Any]]:
    user = await _ai_user(request)
    return await db.ai_schedules.find({"user_id": user["id"]}, {"_id": 0}).sort("created_at", -1).to_list(100)


@router.post("/schedules")
async def create_schedule(payload: Dict[str, Any], request: Request) -> Dict[str, Any]:
    from core_utils import new_id, now_iso
    from services import ai_schedules
    user = await _ai_user(request)
    ctx = await entity_ctx(request)
    if await db.ai_schedules.count_documents({"user_id": user["id"]}) >= 20:
        raise HTTPException(status_code=400, detail="Maksimal 20 jadwal per pengguna")
    try:
        body = ai_schedules.validate(payload, user.get("role") or "")
    except ai_schedules.ScheduleError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    doc = {"id": new_id("aisch"), "user_id": user["id"], "user_name": user.get("name"), "entity_id": ctx.active_entity_id,
           "view_all": bool(ctx.view_all), "active": True, "last_run_at": "", "last_status": "", "created_at": now_iso(), **body}
    await db.ai_schedules.insert_one(dict(doc))
    return doc


@router.patch("/schedules/{schedule_id}")
async def update_schedule(schedule_id: str, payload: Dict[str, Any], request: Request) -> Dict[str, Any]:
    from services import ai_schedules
    user = await _ai_user(request)
    cur = await _own_schedule(schedule_id, user)
    try:
        body = ai_schedules.validate({**cur, **payload, "template_id": cur["template_id"]}, user.get("role") or "")
    except ai_schedules.ScheduleError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    if "active" in payload:
        body["active"] = bool(payload["active"])
    await db.ai_schedules.update_one({"id": schedule_id}, {"$set": body})
    return {**cur, **body}


@router.delete("/schedules/{schedule_id}")
async def delete_schedule(schedule_id: str, request: Request) -> Dict[str, Any]:
    user = await _ai_user(request)
    await _own_schedule(schedule_id, user)
    await db.ai_schedules.delete_one({"id": schedule_id})
    return {"ok": True}


@router.post("/schedules/{schedule_id}/run")
async def run_schedule_now(schedule_id: str, request: Request) -> Dict[str, Any]:
    from services import ai_schedules
    user = await _ai_user(request)
    return await ai_schedules.run_schedule(await _own_schedule(schedule_id, user), trigger="manual")


@router.get("/schedule-runs")
async def list_schedule_runs(request: Request, schedule_id: str = "") -> List[Dict[str, Any]]:
    user = await _ai_user(request)
    q: Dict[str, Any] = {"user_id": user["id"], **({"schedule_id": schedule_id} if schedule_id else {})}
    return await db.ai_schedule_runs.find(q, {"_id": 0, "id": 1, "schedule_id": 1, "title": 1, "status": 1, "error": 1,
                                              "trigger": 1, "created_at": 1}).sort("created_at", -1).to_list(30)


@router.get("/schedule-runs/{run_id}")
async def get_schedule_run(run_id: str, request: Request) -> Dict[str, Any]:
    user = await _ai_user(request)
    q = {"id": run_id} if user.get("role") == "admin" else {"id": run_id, "user_id": user["id"]}
    run = await db.ai_schedule_runs.find_one(q, {"_id": 0})
    if not run:
        raise HTTPException(status_code=404, detail="Riwayat laporan tidak ditemukan")
    owner = {"id": run["user_id"], "role": "admin" if user.get("role") == "admin" else user.get("role")}
    run["results"] = {r["result_id"]: r for r in await _owned_results(run.get("result_ids") or [], owner)}
    return run


@router.get("/snapshots/status")
async def snapshots_status(request: Request) -> Dict[str, Any]:
    from services import analytics_snapshots
    await require_role(request, ["admin", "manager"])
    return await analytics_snapshots.status()


@router.post("/snapshots/run")
async def snapshots_run(request: Request) -> Dict[str, Any]:
    from services import analytics_snapshots
    actor = await require_role(request, ["admin"])
    res = await analytics_snapshots.job_snapshot_daily()
    await audit(actor.get("name", ""), "ai_snapshot_run", "fact_stock_daily", "daily", res)
    return res


@router.get("/usage")
async def usage_summary(request: Request) -> Dict[str, Any]:
    from services import ai_cost
    from services.analytics_time import now_wib
    await require_role(request, ["admin", "manager"])
    start = now_wib().replace(day=1, hour=0, minute=0, second=0, microsecond=0).isoformat()
    rows = await db.ai_usage_log.aggregate([
        {"$match": {"feature": {"$regex": "^bi_"}, "created_at": {"$gte": start}, "demo": {"$ne": True}}},
        {"$group": {"_id": {"feature": "$feature", "model": "$model"}, "calls": {"$sum": 1},
                    "usd": {"$sum": {"$ifNull": ["$cost_usd", 0]}}, "input": {"$sum": {"$ifNull": ["$usage.input_tokens", 0]}},
                    "cached": {"$sum": {"$ifNull": ["$usage.cached_tokens", 0]}}, "output": {"$sum": {"$ifNull": ["$usage.output_tokens", 0]}}}}]).to_list(50)
    items = [{**r["_id"], **{k: r[k] for k in ("calls", "usd", "input", "cached", "output")}} for r in rows]
    tot_in = sum(i["input"] for i in items)
    return {"budget": await ai_cost.budget_state(), "items": items,
            "cache_hit_rate": round(sum(i["cached"] for i in items) / tot_in * 100, 1) if tot_in else None}



@router.get("/usage/daily")
async def usage_daily(request: Request, days: int = 30, scope: str = "bi") -> Dict[str, Any]:
    """Halaman Biaya AI: owner/admin + peran yang boleh memakai asisten (`ai.enabled_roles`)."""
    from services import ai_usage_report
    user = await _ai_user(request)
    # KN-E44 — selain admin/manager hanya melihat pemakaiannya sendiri.
    only = None if user.get("role") in ("admin", "manager") else user["id"]
    return await ai_usage_report.daily_report(days, scope, only_user_id=only)
