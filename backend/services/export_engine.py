"""L11 LN-110/LN-113 — mesin Pusat Export: validasi kolom/filter, kueri ber-scope, join to-one,
preview dan penulisan file (XLSX berpartisi / CSV) yang aman dari formula injection."""
import csv
import re
from datetime import date, timedelta
from typing import Any, Dict, List, Optional, Tuple

from fastapi import HTTPException

import role_registry as reg
from core_utils import COST_SAFE_ROLES, now_iso
from db import db
from dependencies import permission_matrix
from entity_scope import field_for, resolve_list_scope
from services import sales_ownership
from services.customer_service import scope_query as customer_scope
from services.export_registry import DATASET_BY_ID
from services.field_visibility import ensure_product_grades, hidden_for_user, redact
from role_registry import role_in  # peran acuan dikenali (keputusan owner 2026-10-21)

CMP = {"eq", "ne", "contains", "not_contains", "starts", "gt", "gte", "lt", "lte", "in", "empty", "not_empty"}
CMP_LABEL = {"eq": "=", "ne": "≠", "contains": "berisi", "not_contains": "tidak berisi", "starts": "diawali",
             "gt": ">", "gte": "≥", "lt": "<", "lte": "≤", "in": "salah satu dari", "empty": "kosong",
             "not_empty": "tidak kosong"}
MAX_CONDITIONS, MAX_DEPTH = 40, 4
MAX_EXPORT_ROWS = 2_000_000
EXCEL_MAX_ROWS = 1_000_000
FORMULA_PREFIX = ("=", "+", "-", "@", "\t", "\r")
DESIGN_FULL_VIEW = ("admin", "manager", "sales_admin", "md")


def _err(status: int, code: str, message: str, **extra) -> HTTPException:
    return HTTPException(status_code=status, detail={"code": code, "message": message, **extra})


async def _perms(user: Dict[str, Any]) -> Dict[str, List[str]]:
    return (await permission_matrix()).get(user.get("role"), {}) or {}


def _has(perms: Dict[str, List[str]], perm: Optional[tuple]) -> bool:
    if not perm:
        return True
    acts = perms.get(perm[0], []) or []
    return perm[1] in acts or "*" in acts


def _col_denied(col: Dict[str, Any], user, perms, hidden) -> Optional[str]:
    if col.get("vis") and col["vis"] in hidden:
        return f"Atribut {col['vis']} disembunyikan untuk peran Anda"
    if col.get("cost") and not reg.role_in(user.get("role"), COST_SAFE_ROLES):
        return "Kolom biaya (HPP) hanya untuk admin/manajer"
    if col.get("perm") and not _has(perms, col["perm"]):
        return f"Butuh izin {col['perm'][0]}.{col['perm'][1]}"
    return None


async def dataset_access(user: Dict[str, Any], ds: Dict[str, Any]) -> Optional[str]:
    perms = await _perms(user)
    if not _has(perms, ds["perm"]):
        return f"Butuh izin {ds['perm'][0]}.{ds['perm'][1]}"
    if ds.get("vis") and ds["vis"] in await hidden_for_user(user):
        return f"Atribut {ds['vis']} disembunyikan untuk peran Anda"
    return None


async def resolve_columns(user: Dict[str, Any], ds: Dict[str, Any]) -> Tuple[List[Dict[str, Any]], List[Dict[str, Any]]]:
    """Kolom yang boleh dipakai user (dasar + relasi) dan daftar kolom yang ditolak beserta alasan."""
    perms, hidden = await _perms(user), await hidden_for_user(user)
    allowed, denied = [], []
    for c in ds["columns"]:
        why = _col_denied(c, user, perms, hidden)
        row = {**c, "source": "base", "perm": None}
        (denied if why else allowed).append({**row, "reason": why} if why else row)
    for j in ds.get("joins", []):
        jdeny = None if _has(perms, j.get("perm")) else f"Butuh izin {j['perm'][0]}.{j['perm'][1]}"
        for c in j["fields"]:
            why = jdeny or _col_denied(c, user, perms, hidden)
            row = {**c, "key": f"{j['id']}.{c['key']}", "label": c["label"], "source": j["id"],
                   "source_label": j["label"], "filterable": False, "perm": None}
            (denied if why else allowed).append({**row, "reason": why} if why else row)
    return allowed, denied


def _all_keys(ds: Dict[str, Any]) -> set:
    return {c["key"] for c in ds["columns"]} | {f"{j['id']}.{c['key']}" for j in ds.get("joins", []) for c in j["fields"]}


def _coerce(col: Dict[str, Any], v: Any) -> Any:
    t = col["type"]
    if t in ("number", "money"):
        try:
            return float(v)
        except (TypeError, ValueError):
            raise _err(422, "FILTER_VALUE", f"Nilai “{v}” untuk {col['label']} harus angka")
    if t == "bool":
        return v in (True, "true", "1", 1, "ya", "Ya")
    return str(v if v is not None else "")


def compile_filter(node: Optional[Dict[str, Any]], cols: Dict[str, Dict[str, Any]], depth: int = 0,
                   counter: Optional[List[int]] = None) -> Optional[Dict[str, Any]]:
    if not node:
        return None
    counter = counter if counter is not None else [0]
    if depth > MAX_DEPTH:
        raise _err(422, "FILTER_TOO_DEEP", f"Kurung filter maksimal {MAX_DEPTH} tingkat")
    if "items" in node:
        op = node.get("op", "and")
        if op not in ("and", "or"):
            raise _err(422, "FILTER_OP", "Operator grup harus DAN/ATAU")
        parts = [p for p in (compile_filter(i, cols, depth + 1, counter) for i in node.get("items") or []) if p]
        if not parts:
            return None
        return parts[0] if len(parts) == 1 else {f"${op}": parts}
    counter[0] += 1
    if counter[0] > MAX_CONDITIONS:
        raise _err(422, "FILTER_TOO_MANY", f"Maksimal {MAX_CONDITIONS} kondisi")
    key, cmp, val = node.get("field"), node.get("cmp", "eq"), node.get("value")
    col = cols.get(key)
    if not col:
        raise _err(403, "FIELD_NOT_ALLOWED", f"Kolom filter “{key}” tidak tersedia untuk Anda", fields=[key])
    if not col.get("filterable", True):
        raise _err(422, "FILTER_JOIN", f"Filter hanya untuk kolom dataset utama, bukan relasi ({col['label']})")
    if cmp not in CMP:
        raise _err(422, "FILTER_CMP", f"Pembanding “{cmp}” tidak dikenal")
    p = col["path"]
    if cmp == "empty":
        return {p: {"$in": [None, ""]}}
    if cmp == "not_empty":
        return {p: {"$nin": [None, ""]}}
    if cmp in ("contains", "not_contains", "starts"):
        rx = ("^" if cmp == "starts" else "") + re.escape(str(val or ""))
        cond = {"$regex": rx, "$options": "i"}
        return {p: {"$not": re.compile(re.escape(str(val or "")), re.I)}} if cmp == "not_contains" else {p: cond}
    if cmp == "in":
        vals = val if isinstance(val, list) else [x.strip() for x in str(val or "").split(",") if x.strip()]
        return {p: {"$in": [_coerce(col, x) for x in vals]}}
    v = _coerce(col, val)
    return {p: v} if cmp == "eq" else {p: {f"${cmp}": v}}


def describe_filter(node: Optional[Dict[str, Any]], labels: Dict[str, str], top: bool = True) -> str:
    if not node:
        return "(tanpa filter)" if top else ""
    if "items" in node:
        parts = [s for s in (describe_filter(i, labels, False) for i in node.get("items") or []) if s]
        joined = f" {'DAN' if node.get('op', 'and') == 'and' else 'ATAU'} ".join(parts)
        return joined if top or len(parts) < 2 else f"({joined})"
    cmp = node.get("cmp", "eq")
    val = "" if cmp in ("empty", "not_empty") else f" “{node.get('value')}”"
    return f"{labels.get(node.get('field'), node.get('field'))} {CMP_LABEL.get(cmp, cmp)}{val}"


def _date_bounds(col: Dict[str, Any], dfrom: str, dto: str) -> Dict[str, str]:
    """Batas inklusif berbasis tanggal WIB (UTC+7). Kolom datetime tersimpan ISO UTC."""
    out: Dict[str, str] = {}
    try:
        f = date.fromisoformat(dfrom) if dfrom else None
        t = date.fromisoformat(dto) if dto else None
    except ValueError:
        raise _err(422, "DATE_FORMAT", "Format tanggal harus YYYY-MM-DD")
    if f and t and f > t:
        raise _err(422, "DATE_RANGE", "Tanggal awal setelah tanggal akhir")
    if col["type"] == "datetime":
        if f:
            out["$gte"] = f"{(f - timedelta(days=1)).isoformat()}T17:00:00"
        if t:
            out["$lt"] = f"{t.isoformat()}T17:00:00"
    else:
        if f:
            out["$gte"] = f.isoformat()
        if t:
            out["$lt"] = (t + timedelta(days=1)).isoformat()
    return out


async def _owner_clause(ds: Dict[str, Any], user: Dict[str, Any]) -> Optional[Dict[str, Any]]:
    mode, uid = ds.get("owner_scope"), str(user.get("id") or "")
    locked = reg.role_in(user.get("role"), sales_ownership.OWNER_LOCKED_ROLES)
    if mode == "sales_order" and locked:
        return sales_ownership.owner_clause(user)
    if mode == "created_by" and locked:
        return {"created_by": uid}
    if mode == "customer" and locked:
        return {"assigned_sales_id": uid} if not role_in(user.get("role"), ("sales",)) else customer_scope(user, {})
    if mode == "via_order" and locked:
        ids = await db.sales_orders.distinct("id", sales_ownership.owner_clause(user))
        return {"order_id": {"$in": ids}}
    if mode == "design" and not reg.role_in(user.get("role"), DESIGN_FULL_VIEW):
        return {"assigned_to": uid}
    return None


async def build_plan(spec: Dict[str, Any], user: Dict[str, Any], ctx) -> Dict[str, Any]:
    ds = DATASET_BY_ID.get(spec.get("dataset") or "")
    if not ds:
        raise _err(404, "DATASET_UNKNOWN", "Dataset tidak dikenal")
    why = await dataset_access(user, ds)
    if why:
        raise _err(403, "DATASET_NOT_ALLOWED", f"Dataset “{ds['label']}” tidak boleh Anda akses: {why}")
    allowed, denied = await resolve_columns(user, ds)
    by_key = {c["key"]: c for c in allowed}
    keys = list(dict.fromkeys(spec.get("columns") or [c["key"] for c in allowed if c["source"] == "base"]))
    known = _all_keys(ds)
    unknown = [k for k in keys if k not in known]
    if unknown:
        raise _err(422, "FIELD_UNKNOWN", f"Kolom tidak dikenal: {', '.join(unknown)}", fields=unknown)
    blocked = [k for k in keys if k not in by_key]
    if blocked:
        raise _err(403, "FIELD_NOT_ALLOWED", "Kolom tidak diizinkan untuk peran Anda", fields=blocked)
    if not keys:
        raise _err(422, "NO_COLUMNS", "Pilih minimal satu kolom")
    base_cols = {k: c for k, c in by_key.items() if c["source"] == "base"}
    post = compile_filter(spec.get("filter"), base_cols)
    pre: Dict[str, Any] = {}
    if field_for(ds["collection"]) is not None:
        pre = resolve_list_scope(ds["collection"], {}, ctx, "all" if spec.get("entity_scope") == "all" else None)
    ands = [pre] if pre else []
    owner = await _owner_clause(ds, user)
    if owner:
        ands.append(owner)
    dkey = spec.get("date_field")
    date_label = ""
    if dkey and (spec.get("date_from") or spec.get("date_to")):
        dcol = next((d for d in ds["date_fields"] if d["key"] == dkey), None)
        if not dcol:
            raise _err(422, "DATE_FIELD", "Basis tanggal tidak tersedia pada dataset ini")
        ands.append({dcol["path"]: _date_bounds(dcol, spec.get("date_from") or "", spec.get("date_to") or "")})
        date_label = f"{dcol['label']}: {spec.get('date_from') or '…'} s/d {spec.get('date_to') or '…'} (WIB, inklusif)"
    sort = []
    for s in spec.get("sort") or []:
        col = base_cols.get(s.get("key"))
        if not col:
            raise _err(403, "FIELD_NOT_ALLOWED", f"Kolom urut “{s.get('key')}” tidak tersedia", fields=[s.get("key")])
        sort.append((col["path"], -1 if s.get("dir") == "desc" else 1))
    labels = {c["key"]: c["label"] for c in allowed}
    pipeline: List[Dict[str, Any]] = []
    if ands:
        pipeline.append({"$match": {"$and": ands} if len(ands) > 1 else ands[0]})
    if ds.get("unwind"):
        pipeline.append({"$unwind": {"path": f"${ds['unwind']}", "includeArrayIndex": "_line",
                                     "preserveNullAndEmptyArrays": bool(ds.get("unwind_preserve"))}})
    pipeline += ds.get("stages") or []
    if post:
        pipeline.append({"$match": post})
    sort_doc = {p: d for p, d in sort}
    sort_doc.setdefault(ds["key"], 1)
    if ds.get("unwind"):
        sort_doc.setdefault("_line", 1)
    cols = [by_key[k] for k in keys]
    group = None
    if spec.get("group_by"):
        cols, group = _group_plan(ds, spec["group_by"], cols, base_cols)
    return {"ds": ds, "columns": cols, "group": group, "pipeline": pipeline, "sort": sort_doc, "hidden": await hidden_for_user(user),
            "filter_text": describe_filter(spec.get("filter"), labels), "date_text": date_label or "(semua tanggal)",
            "sort_text": ", ".join(f"{labels.get(s.get('key'))} {'↓' if s.get('dir') == 'desc' else '↑'}"
                                   for s in spec.get("sort") or []) or "(bawaan: kunci dataset)",
            "entity_ids": list(ctx.allowed_entity_ids if spec.get("entity_scope") == "all" or ctx.view_all
                               else [ctx.active_entity_id]),
            "denied": [{"key": d["key"], "label": d["label"], "reason": d["reason"]} for d in denied]}


ROWS_COL = {"key": "_rows", "label": "Jumlah baris", "type": "number", "path": "_rows", "source": "base",
            "derived": "group_count"}
#: AUD-EXP-01 — kolom TURUNAN yang ditambahkan engine (bukan kolom dataset): asal & aturan izinnya. `_rows`
#: = hitungan baris per kelompok; tidak membawa nilai kolom apa pun, jadi cukup hak atas dataset + kolom
#: pengelompokan (yang tetap diperiksa ulang saat unduh). Jangan dihapus dari pemeriksaan unduh.
DERIVED_COLUMNS = {ROWS_COL["key"]: {"origin": "group_count", "requires": "dataset"}}


def _group_plan(ds, group_by, cols, base_cols):
    """LN-111 — ringkasan per kunci: angka diagregasi (sum/avg) dari baris grain yang sama, tanpa join."""
    if len(group_by) > 3:
        raise _err(422, "GROUP_TOO_MANY", "Kelompokkan maksimal 3 kolom")
    bad = [g for g in group_by if g not in base_cols]
    if bad:
        raise _err(403, "FIELD_NOT_ALLOWED", "Kolom pengelompokan tidak tersedia", fields=bad)
    if any(c["source"] != "base" for c in cols):
        raise _err(422, "GROUP_JOIN", "Mode ringkasan hanya untuk kolom dataset utama (tanpa relasi)")
    keys = [base_cols[g] for g in group_by]
    measures = [c for c in cols if c["type"] in ("number", "money") and c["key"] not in group_by]
    agg = ds.get("agg") or {}
    measures = [c for c in measures if agg.get(c["key"]) != "none"]
    measures = [{**c, "agg": agg.get(c["key"], "sum"),
                 "label": c["label"] + (" (rata-rata, kosong diabaikan)" if agg.get(c["key"]) == "avg" else "")}
                for c in measures]
    group = {"_id": {g["key"]: f"${g['path']}" for g in keys}, "_rows": {"$sum": 1},
             **{m["key"]: {f"${m['agg']}": f"${m['path']}"} for m in measures}}
    return keys + measures + [ROWS_COL], {"stage": group, "keys": [g["key"] for g in keys]}


def _grouped(plan):
    g = plan["group"]
    return plan["pipeline"] + [{"$group": g["stage"]}]


async def count(plan: Dict[str, Any]) -> int:
    pipe = _grouped(plan) if plan.get("group") else plan["pipeline"]
    res = await db[plan["ds"]["collection"]].aggregate(pipe + [{"$count": "n"}]).to_list(1)
    return res[0]["n"] if res else 0


def _get(doc: Any, path: str) -> Any:
    for part in path.split("."):
        if isinstance(doc, dict):
            doc = doc.get(part)
        else:
            return None
    return doc


def _norm(col: Dict[str, Any], v: Any) -> Any:
    if v is None or v == "":
        return None
    if col["type"] in ("number", "money"):
        try:
            return float(v)
        except (TypeError, ValueError):
            return None
    if col["type"] == "bool":
        return bool(v)
    if isinstance(v, list):
        return ", ".join(str(x) for x in v if not isinstance(x, (dict, list)))
    if isinstance(v, dict):
        return None
    return str(v)


def _shape_grouped(plan, docs):
    keys = plan["group"]["keys"]
    rows = [{c["key"]: _norm(c, d["_id"].get(c["key"]) if c["key"] in keys else d.get(c["key"])) for c in plan["columns"]}
            for d in docs]
    return redact(rows, plan["hidden"]) if plan["hidden"] else rows


async def _fetch_grouped(plan, skip, limit):
    keys = plan["group"]["keys"]
    pipe = _grouped(plan) + [{"$sort": {**{f"_id.{k}": 1 for k in keys}}}, {"$skip": skip}, {"$limit": limit}]
    docs = await db[plan["ds"]["collection"]].aggregate(pipe, allowDiskUse=True).to_list(limit)
    return _shape_grouped(plan, docs)


def _used_joins(plan) -> Dict[str, Dict[str, Any]]:
    joins = {j["id"]: j for j in plan["ds"].get("joins", [])}
    return {jid: joins[jid] for jid in {c["source"] for c in plan["columns"] if c["source"] != "base"}}


def _base_proj(plan) -> Dict[str, Any]:
    ds = plan["ds"]
    proj = {"_id": 0, ds["key"]: 1}
    for c in plan["columns"]:
        if c["source"] == "base":
            proj[c["path"]] = 1
    for j in _used_joins(plan).values():
        proj[j["local"]] = 1
    return proj


async def _lookup_maps(plan, ids_by_join: Dict[str, List[Any]]) -> Dict[str, Dict[str, Dict[str, Any]]]:
    lookups: Dict[str, Dict[str, Dict[str, Any]]] = {}
    for jid, j in _used_joins(plan).items():
        ids = ids_by_join.get(jid) or []
        fp = {"_id": 0, j["foreign"]: 1, **{f["path"]: 1 for f in j["fields"]}}
        found: List[Dict[str, Any]] = []
        for i in range(0, len(ids), 5000):
            chunk = ids[i:i + 5000]
            found += await db[j["collection"]].find({j["foreign"]: {"$in": chunk}}, fp).to_list(len(chunk) + 1)
        lookups[jid] = {str(f.get(j["foreign"])): f for f in found}
    return lookups


async def _shape(plan, docs, lookups) -> List[Dict[str, Any]]:
    joins = _used_joins(plan)
    rows = []
    for d in docs:
        row = {}
        for c in plan["columns"]:
            if c["source"] == "base":
                row[c["key"]] = _norm(c, _get(d, c["path"]))
            else:
                j = joins[c["source"]]
                ref = lookups[c["source"]].get(str(_get(d, j["local"]) or ""), {})
                row[c["key"]] = _norm(c, _get(ref, c["path"]))
        rows.append(row)
    if plan["hidden"]:
        await ensure_product_grades()
        rows = redact(rows, plan["hidden"])
    return rows


async def fetch(plan: Dict[str, Any], skip: int, limit: int) -> List[Dict[str, Any]]:
    """Baca LIVE per halaman — untuk preview. File export memakai snapshot (`materialize`), bukan ini."""
    if plan.get("group"):
        return await _fetch_grouped(plan, skip, limit)
    ds = plan["ds"]
    pipe = plan["pipeline"] + [{"$sort": plan["sort"]}, {"$skip": skip}, {"$limit": limit}, {"$project": _base_proj(plan)}]
    docs = await db[ds["collection"]].aggregate(pipe, allowDiskUse=True).to_list(limit)
    ids = {jid: list({v for d in docs for v in [_get(d, j["local"])] if v and not isinstance(v, (dict, list))})
           for jid, j in _used_joins(plan).items()}
    return await _shape(plan, docs, await _lookup_maps(plan, ids))


# ─── AUD-EXP-02 — konsistensi export: SATU pembacaan sumber (snapshot) ──────────────────────────────────────
# Tanpa replica set (keputusan T-01) tidak ada snapshot read Mongo. Definisi konsistensi export: himpunan baris =
# hasil SATU agregasi atas sumber yang dimaterialisasi ke koleksi sementara `export_snap_<job>` (unik per dokumen
# sumber + baris unwind). File dibaca dari snapshot statis itu lewat SATU cursor terurut; jumlah baris, ringkasan &
# nilai join juga dari snapshot itu — perubahan sumber selama penulisan file (insert/hapus/ubah filter/urutan/
# nilai) tidak menggeser, menggandakan, atau mencampur versi baris. Waktu snapshot dicatat di sheet Ringkasan.
SNAP_PREFIX = "export_snap_"


def _clean_proj(paths: List[str]) -> Dict[str, int]:
    """Proyeksi tanpa tabrakan jalur (`a` dan `a.b` → cukup `a`)."""
    keep: List[str] = []
    for path in sorted(set(paths), key=lambda x: (x.count("."), x)):
        if not any(path == k or path.startswith(k + ".") for k in keep):
            keep.append(path)
    return {k: 1 for k in keep}


def _snap_paths(plan) -> List[str]:
    ds = plan["ds"]
    paths = [ds["key"], *plan["sort"].keys()]
    paths += [k for k in _base_proj(plan) if k != "_id"]
    sp = plan.get("summary_plan")
    if sp:   # ringkasan dihitung dari snapshot yang sama → jalur kunci & ukurannya ikut disimpan
        paths += [c["path"] for c in sp["columns"] if c.get("source") == "base" and c["key"] != ROWS_COL["key"]]
    return paths


async def materialize(plan: Dict[str, Any], name: str) -> int:
    await db[name].drop()
    src = db[plan["ds"]["collection"]]
    if plan.get("group"):
        pipe = _grouped(plan)
    else:
        ds = plan["ds"]
        uniq: Dict[str, Any] = {"o": "$_id"}
        if ds.get("unwind"):
            uniq["l"] = "$_line"
        pipe = plan["pipeline"] + [
            {"$project": {**_clean_proj(_snap_paths(plan)), "_id": 1}},
            {"$group": {"_id": uniq, "d": {"$first": "$$ROOT"}}},   # satu versi per dokumen sumber (+ baris)
            {"$replaceRoot": {"newRoot": "$d"}}, {"$unset": "_id"}]
    await src.aggregate(pipe + [{"$out": name}], allowDiskUse=True).to_list(None)
    return await db[name].count_documents({})


def _snap_sort(plan) -> Dict[str, int]:
    return {f"_id.{k}": 1 for k in plan["group"]["keys"]} if plan.get("group") else plan["sort"]


async def _snap_batches(name: str, batch: int, plan: Dict[str, Any]):
    """Satu cursor terurut atas snapshot statis (tanpa skip/limit berulang) → tidak ada pergeseran halaman."""
    pipe: List[Dict[str, Any]] = [{"$sort": _snap_sort(plan)}]
    if not plan.get("group"):
        pipe.append({"$project": {"_id": 0}})
    chunk: List[Dict[str, Any]] = []
    async for d in db[name].aggregate(pipe, allowDiskUse=True, batchSize=batch):
        chunk.append(d)
        if len(chunk) >= batch:
            yield chunk
            chunk = []
    if chunk:
        yield chunk


async def _snap_lookups(plan, name: str):
    """Nilai join dibaca SEKALI untuk seluruh snapshot — baris dengan kunci relasi sama memakai versi yang sama."""
    ids: Dict[str, List[Any]] = {}
    for jid, j in _used_joins(plan).items():
        vals = await db[name].aggregate([{"$group": {"_id": f"${j['local']}"}}], allowDiskUse=True).to_list(None)
        ids[jid] = [v["_id"] for v in vals if v["_id"] and not isinstance(v["_id"], (dict, list))]
    return await _lookup_maps(plan, ids)


async def _snap_summary(sp, name: str) -> List[Dict[str, Any]]:
    keys = sp["group"]["keys"]
    pipe = [{"$group": sp["group"]["stage"]}, {"$sort": {f"_id.{k}": 1 for k in keys}}]
    return _shape_grouped(sp, await db[name].aggregate(pipe, allowDiskUse=True).to_list(None))


async def sweep_snapshots() -> None:
    """Snapshot yatim (proses mati di tengah) dibuang bila job-nya tidak sedang berjalan."""
    for n in await db.list_collection_names(filter={"name": {"$regex": f"^{SNAP_PREFIX}"}}):
        job = await db.export_jobs.find_one({"id": n[len(SNAP_PREFIX):]}, {"_id": 0, "status": 1})
        if not job or job.get("status") not in ("queued", "running"):
            await db[n].drop()


def safe_text(v: Any) -> Any:
    """LN-113-T03 — teks diawali karakter formula diberi apostrof agar dibaca sebagai teks."""
    if isinstance(v, str) and v.startswith(FORMULA_PREFIX):
        return "'" + v
    return v


def summary_lines(job: Dict[str, Any], plan: Dict[str, Any]) -> List[List[Any]]:
    ds = plan["ds"]
    lines = [["Dataset", ds["label"]], ["Grain", ds["grain"]], ["Kunci", ds["key"]],
             ["Dibuat", job["created_at"]], ["Oleh", f"{job['user_name']} ({job['role']})"],
             ["Badan usaha", ", ".join(plan["entity_ids"])], ["Periode", plan["date_text"]],
             ["Filter", plan["filter_text"]], ["Urutan", plan["sort_text"]],
             ["Total baris", job.get("written_rows", 0)], ["Format", job["format"].upper()]]
    if job.get("snapshot_at"):
        lines.append(["Data per", f"{job['snapshot_at']} (snapshot sekali baca; baris unik per dokumen sumber)"])
    for p in job.get("partitions") or []:
        lines.append([f"Partisi {p['sheet']}", f"baris {p['from']}–{p['to']}"])
    if plan["denied"]:
        lines.append(["Kolom tidak tersedia untuk peran ini", ", ".join(d["label"] for d in plan["denied"])])
    lines.append([])
    lines.append(["Definisi kolom", "Label", "Tipe", "Sumber"])
    for c in plan["columns"]:
        lines.append([c["key"], c["label"], c["type"], c.get("source_label") or ds["label"]])
    return lines


async def write_file(job: Dict[str, Any], plan: Dict[str, Any], path: str, rows_per_sheet: int,
                     batch: int = 2000) -> Dict[str, Any]:
    """AUD-EXP-02 — sumber dibaca SEKALI ke snapshot; file, ringkasan & join dari snapshot itu."""
    name = f"{SNAP_PREFIX}{job['id']}"
    try:
        await sweep_snapshots()
    except Exception:  # noqa: BLE001 — pembersihan oportunistik, tidak menggagalkan export
        pass
    snapshot_at = now_iso()
    try:
        await materialize(plan, name)
        grouped = bool(plan.get("group"))
        lookups = {} if grouped else await _snap_lookups(plan, name)

        async def rows_iter():
            async for docs in _snap_batches(name, batch, plan):
                yield _shape_grouped(plan, docs) if grouped else await _shape(plan, docs, lookups)

        res = await _write(job, plan, path, rows_per_sheet, rows_iter(), name, snapshot_at)
    finally:
        await db[name].drop()
    return {**res, "snapshot_at": snapshot_at}


async def _write(job, plan, path, rows_per_sheet, rows_iter, name, snapshot_at) -> Dict[str, Any]:
    cols = plan["columns"]
    header = [c["label"] for c in cols]
    written, partitions = 0, []
    if job["format"] == "csv":
        with open(path, "w", newline="", encoding="utf-8-sig") as fh:
            w = csv.writer(fh)
            w.writerow(header)
            async for rows in rows_iter:
                for r in rows:
                    w.writerow([safe_text(r.get(c["key"])) if r.get(c["key"]) is not None else "" for c in cols])
                written += len(rows)
        return {"written_rows": written, "partitions": []}
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    wb = Workbook(write_only=True)
    ws, in_sheet, sheet_no = None, 0, 0

    def new_sheet():
        nonlocal ws, in_sheet, sheet_no
        sheet_no += 1
        ws = wb.create_sheet(f"Data {sheet_no}")
        ws.append(header)
        in_sheet = 0
        partitions.append({"sheet": f"Data {sheet_no}", "from": written + 1, "to": written})

    new_sheet()
    async for rows in rows_iter:
        for r in rows:
            if in_sheet >= rows_per_sheet:
                new_sheet()
            out = []
            for c in cols:
                v = r.get(c["key"])
                cell = WriteOnlyCell(ws, value=v)
                if isinstance(v, str):
                    cell.data_type = "s"
                out.append(cell)
            ws.append(out)
            in_sheet += 1
            written += 1
            partitions[-1]["to"] = written
    job = {**job, "written_rows": written, "partitions": partitions, "snapshot_at": snapshot_at}
    sp = plan.get("summary_plan")
    if sp:
        label = "+".join(c["label"] for c in sp["columns"] if c["key"] in sp["group"]["keys"])
        ws2 = wb.create_sheet(f"Ringkasan per {label}"[:31])
        ws2.append([c["label"] for c in sp["columns"]])
        for r in await _snap_summary(sp, name):
            ws2.append([WriteOnlyCell(ws2, value=r.get(c["key"])) for c in sp["columns"]])
    summary = wb.create_sheet("Ringkasan")
    for line in summary_lines(job, plan):
        cells = []
        for v in line:
            cell = WriteOnlyCell(summary, value=v)
            if isinstance(v, str):
                cell.data_type = "s"
            cells.append(cell)
        summary.append(cells)
    wb._sheets.insert(0, wb._sheets.pop())
    wb.save(path)
    return {"written_rows": written, "partitions": partitions}
