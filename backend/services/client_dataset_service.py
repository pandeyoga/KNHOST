"""L01 LN-011 — IMPOR DATASET CLIENT (DATA-01): artikel · tahap · jenis kain · supplier/warna.

Prinsip (paket: "jangan buat mapping bisnis fiktif"):
  * Hanya MEMPERBARUI master yang cocok (template via kode artikel/nama, SKU via kode SKU atau
    artikel+warna internal) dan memetakan supplier/warna supplier (`supplier_items`).
  * TIDAK membuat artikel/SKU/supplier baru — baris tak cocok dilaporkan untuk diputuskan owner.
  * Nilai tahap/jenis kain di luar master dilaporkan (bukan ditebak).
  * Selalu dry-run dulu; commit menyimpan nilai SEBELUM per dokumen → bisa dibatalkan (rollback).
"""
import csv
import io
import re
from typing import Any, Dict, List, Optional, Tuple

from fastapi import HTTPException

import domain_registry as dr
from core_utils import new_id, now_iso
from db import db

COLS = {
    "kode_artikel": ("kode_artikel", "kode artikel", "sku_prefix", "artikel_kode", "article_code"),
    "nama_artikel": ("nama_artikel", "nama artikel", "artikel", "nama_internal", "article_name"),
    "sku": ("sku", "kode_sku", "kode sku"),
    "tahap": ("tahap", "stage", "tahap_bahan"),
    "jenis_kain": ("jenis_kain", "jenis kain", "fabric_type", "jenis"),
    "kategori": ("kategori", "category"),
    "warna_internal": ("warna_internal", "warna internal", "warna", "color"),
    "supplier": ("supplier", "nama_supplier", "nama supplier", "pemasok"),
    "kode_artikel_supplier": ("kode_artikel_supplier", "supplier_sku", "kode supplier", "kode_supplier_artikel"),
    "nama_artikel_supplier": ("nama_artikel_supplier", "supplier_item_name", "nama artikel supplier"),
    "warna_supplier": ("warna_supplier", "warna supplier", "supplier_color"),
    "kode_warna_supplier": ("kode_warna_supplier", "kode warna supplier", "supplier_color_code"),
    "warna_customer": ("warna_customer", "warna customer", "customer_color"),
}
TEMPLATE_HEADER = list(COLS.keys())
# Sinonim BAKU saja (ejaan/istilah), bukan pemetaan bisnis.
STAGE_SYN = {"greige": "grey", "gray": "grey", "grey/greige": "grey", "jadi": "finished", "benang": "yarn"}
FABRIC_SYN = {"knitting": "knit", "rajut": "knit", "tenun": "woven", "woven fabric": "woven"}


def _norm(s: Any) -> str:
    return " ".join(str(s or "").strip().split())


def parse(content: bytes, filename: str) -> List[Dict[str, str]]:
    if (filename or "").lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        ws = load_workbook(io.BytesIO(content), read_only=True, data_only=True).active
        it = ws.iter_rows(values_only=True)
        header = [_norm(h).lower() for h in next(it, [])]
        raw = [dict(zip(header, ["" if v is None else str(v) for v in r])) for r in it]
    else:
        text = content.decode("utf-8-sig", errors="replace")
        sample = text[:2000]
        delim = ";" if sample.count(";") > sample.count(",") else ","
        raw = [{_norm(k).lower(): v for k, v in r.items()} for r in csv.DictReader(io.StringIO(text), delimiter=delim)]
    alias = {a: k for k, al in COLS.items() for a in al}
    rows = []
    for r in raw:
        row = {alias[k]: _norm(v) for k, v in r.items() if k in alias}
        if any(row.values()):
            rows.append(row)
    if not rows:
        raise HTTPException(status_code=400, detail="File kosong atau kolom tidak dikenali. Unduh template.")
    if len(rows) > 5000:
        raise HTTPException(status_code=400, detail="Maksimal 5.000 baris per impor.")
    return rows


async def _indexes() -> Tuple[Dict[str, Any], Dict[str, Any], Dict[str, Any], Dict[str, Any]]:
    tpls = await db.product_templates.find({}, {"_id": 0, "id": 1, "name": 1, "sku_prefix": 1, "stage": 1,
                                                 "fabric_type": 1, "category": 1}).to_list(20000)
    by_code = {t["sku_prefix"].upper(): t for t in tpls if t.get("sku_prefix")}
    by_name = {t["name"].lower(): t for t in tpls if t.get("name")}
    prods = await db.products.find({}, {"_id": 0, "id": 1, "sku": 1, "name": 1, "template_id": 1, "color": 1,
                                        "stage": 1, "fabric_type": 1, "customer_color": 1}).to_list(50000)
    by_sku = {p["sku"].upper(): p for p in prods if p.get("sku")}
    by_tpl_color: Dict[str, List[Dict[str, Any]]] = {}
    for p in prods:
        by_tpl_color.setdefault(f"{p.get('template_id')}|{str(p.get('color') or '').lower()}", []).append(p)
    sups = await db.suppliers.find({}, {"_id": 0, "id": 1, "name": 1, "code": 1}).to_list(5000)
    sup: Dict[str, Any] = {}
    for s in sups:
        for key in {(s.get("name") or "").lower(), (s.get("code") or "").lower()} - {""}:
            sup[key] = "AMBIGUOUS" if key in sup and sup[key]["id"] != s["id"] else s
    return {"code": by_code, "name": by_name}, {"sku": by_sku, "tc": by_tpl_color}, sup, {}


async def dry_run(rows: List[Dict[str, str]]) -> Dict[str, Any]:
    tidx, pidx, sup, _ = await _indexes()
    ssku: Dict[Tuple[str, str], str] = {}   # (supplier_id, supplier_sku) → product_id (unik di DB, semua status)
    async for x in db.supplier_items.find({"supplier_sku": {"$nin": [None, ""]}},
                                          {"_id": 0, "supplier_id": 1, "supplier_sku": 1, "product_id": 1, "sku": 1}):
        ssku[(x["supplier_id"], x["supplier_sku"].upper())] = x.get("sku") or x["product_id"]
    pid_sku = {p["id"]: p["sku"] for p in pidx["sku"].values()}
    stages, fabrics = set(dr.values_of("stage")), set(dr.values_of("fabric_type"))
    out, unknown = [], {"tahap": {}, "jenis_kain": {}, "supplier": {}}
    plan = {"templates": {}, "products": {}, "supplier_items": []}
    for i, r in enumerate(rows, start=2):
        res: Dict[str, Any] = {"line": i, "input": r, "errors": [], "warnings": [], "changes": []}
        tpl = tidx["code"].get(r.get("kode_artikel", "").upper()) or tidx["name"].get(r.get("nama_artikel", "").lower())
        prods: List[Dict[str, Any]] = []
        if r.get("sku"):
            p = pidx["sku"].get(r["sku"].upper())
            prods = [p] if p else []
            tpl = tpl or ({"id": p["template_id"]} if p and p.get("template_id") else None)
        elif tpl and r.get("warna_internal"):
            prods = pidx["tc"].get(f"{tpl['id']}|{r['warna_internal'].lower()}", [])
        if not tpl and not prods:
            res["errors"].append("Artikel/SKU tidak ditemukan di master — tidak dibuat otomatis (keputusan owner).")
        stage = r.get("tahap", "").lower()
        fabric = r.get("jenis_kain", "").lower()
        for kind, val, syn in (("tahap", stage, STAGE_SYN), ("jenis kain", fabric, FABRIC_SYN)):
            if val in syn:
                res["warnings"].append(f"{kind} '{val}' dibaca sebagai '{syn[val]}' (sinonim baku).")
        stage, fabric = STAGE_SYN.get(stage, stage), FABRIC_SYN.get(fabric, fabric)
        fabric = re.sub(r"[^a-z0-9_]", "_", fabric).strip("_")   # sama dengan normalisasi kode master jenis kain
        if stage and stage not in stages:
            unknown["tahap"][stage] = unknown["tahap"].get(stage, 0) + 1
            res["errors"].append(f"Tahap '{r['tahap']}' tidak ada di master ({', '.join(sorted(stages))}).")
        if fabric and fabric not in fabrics:
            unknown["jenis_kain"][fabric] = unknown["jenis_kain"].get(fabric, 0) + 1
            res["errors"].append(f"Jenis kain '{r['jenis_kain']}' tidak ada di master ({', '.join(sorted(fabrics))}).")
        s = sup.get(r.get("supplier", "").lower()) if r.get("supplier") else None
        if s == "AMBIGUOUS":
            res["errors"].append(f"Nama supplier '{r['supplier']}' dipakai lebih dari satu pemasok — tulis KODE supplier.")
            s = None
        elif r.get("supplier") and not s:
            unknown["supplier"][r["supplier"]] = unknown["supplier"].get(r["supplier"], 0) + 1
            res["errors"].append(f"Supplier '{r['supplier']}' tidak ada di master pemasok.")
        if not res["errors"]:
            if tpl and tpl.get("name") is not None:
                t = plan["templates"].setdefault(tpl["id"], {})
                for fld, val in (("stage", stage), ("fabric_type", fabric), ("category", r.get("kategori", ""))):
                    if val and tpl.get(fld) != val:
                        if t.get(fld) not in (None, val):
                            res["warnings"].append(f"Baris lain menetapkan {fld} artikel berbeda ({t[fld]} vs {val}).")
                        t[fld] = val
                        res["changes"].append(f"artikel {tpl.get('sku_prefix') or tpl['name']}: {fld} {tpl.get(fld) or '—'} → {val}")
            for p in prods:
                pc = plan["products"].setdefault(p["id"], {})
                for fld, val in (("stage", stage), ("fabric_type", fabric), ("customer_color", r.get("warna_customer", ""))):
                    if val and p.get(fld) != val:
                        pc[fld] = val
                        res["changes"].append(f"SKU {p['sku']}: {fld} {p.get(fld) or '—'} → {val}")
                key = (s["id"], r.get("kode_artikel_supplier", "").upper()) if s else None
                if key and key[1] and ssku.get(key) not in (None, p["sku"], p["id"]):
                    res["errors"].append(
                        f"Kode artikel supplier {r['kode_artikel_supplier']} sudah dipakai SKU {ssku[key]} untuk "
                        f"{s['name']} — kode supplier harus unik per SKU (mis. gabungkan dengan kode warna).")
                    continue
                if key and key[1]:
                    ssku[key] = p["sku"]
                if s and (r.get("warna_supplier") or r.get("kode_artikel_supplier") or r.get("kode_warna_supplier")):
                    plan["supplier_items"].append({
                        "product_id": p["id"], "sku": p["sku"], "product_name": p.get("name", ""),
                        "supplier_id": s["id"], "supplier_name": s["name"],
                        "supplier_sku": r.get("kode_artikel_supplier", ""),
                        "supplier_item_name": r.get("nama_artikel_supplier", ""),
                        "supplier_color": r.get("warna_supplier", ""),
                        "supplier_color_code": r.get("kode_warna_supplier", ""),
                        "internal_color_code": r.get("warna_internal", "")})
                    res["changes"].append(f"SKU {p['sku']}: peta supplier {s['name']} warna "
                                          f"{r.get('warna_supplier') or r.get('kode_warna_supplier')}")
            if not res["changes"]:
                res["warnings"].append("Tidak ada perubahan (data sudah sama).")
        if res["errors"]:
            res["changes"] = []
        res["status"] = "error" if res["errors"] else ("ok" if res["changes"] else "same")
        out.append(res)
    plan["templates"] = {k: v for k, v in plan["templates"].items() if v}
    plan["products"] = {k: v for k, v in plan["products"].items() if v}
    summary = {k: sum(1 for x in out if x["status"] == k) for k in ("ok", "same", "error")}
    summary.update(rows=len(out), templates=len(plan["templates"]), products=len(plan["products"]),
                   supplier_items=len(plan["supplier_items"]))
    return {"summary": summary, "rows": out, "unknown_values": unknown, "plan": plan}


async def save_batch(rows: List[Dict[str, str]], report: Dict[str, Any], filename: str, actor: str) -> Dict[str, Any]:
    doc = {"id": new_id("cdi"), "filename": filename, "status": "dry_run", "created_by": actor, "created_at": now_iso(),
           "rows": rows, "summary": report["summary"], "unknown_values": report["unknown_values"],
           "plan": report["plan"]}
    await db.client_dataset_imports.insert_one(dict(doc))
    return {k: v for k, v in doc.items() if k not in ("rows", "plan", "_id")}


async def commit(batch_id: str, actor: str) -> Dict[str, Any]:
    b = await db.client_dataset_imports.find_one_and_update(
        {"id": batch_id, "status": "dry_run"}, {"$set": {"status": "committing", "committed_by": actor}})
    if not b:
        raise HTTPException(status_code=409, detail="Batch tidak ada atau sudah diproses.")
    from pymongo.errors import DuplicateKeyError
    try:
        return await _apply(b, batch_id, actor)
    except Exception as exc:
        await db.client_dataset_imports.update_one({"id": batch_id, "status": "committing"}, {"$set": {
            "status": "failed", "error": str(exc)[:300]}})
        if isinstance(exc, DuplicateKeyError):
            raise HTTPException(status_code=409, detail="Data bentrok dengan master (kode ganda). Impor ditandai "
                                                        "gagal — batalkan lalu cek ulang file.") from exc
        raise


async def _apply(b: Dict[str, Any], batch_id: str, actor: str) -> Dict[str, Any]:
    # dry-run ulang supaya tidak menulis rencana basi bila master berubah sejak upload
    fresh = await dry_run(b["rows"])
    plan, before = fresh["plan"], {"templates": {}, "products": {}, "supplier_items": []}
    async def snap():   # simpan 'before' bertahap — batch gagal di tengah tetap bisa dibatalkan
        await db.client_dataset_imports.update_one({"id": batch_id}, {"$set": {"before": before}})
    for tid, ch in plan["templates"].items():
        cur = await db.product_templates.find_one({"id": tid}, {"_id": 0, **{k: 1 for k in ch}})
        before["templates"][tid] = {k: (cur or {}).get(k) for k in ch}
        await snap()
        await db.product_templates.update_one({"id": tid}, {"$set": {**ch, "updated_at": now_iso(), "import_batch": batch_id}})
    for pid, ch in plan["products"].items():
        cur = await db.products.find_one({"id": pid}, {"_id": 0, **{k: 1 for k in ch}})
        before["products"][pid] = {k: (cur or {}).get(k) for k in ch}
        await snap()
        await db.products.update_one({"id": pid}, {"$set": {**ch, "updated_at": now_iso()}})
    for si in plan["supplier_items"]:
        cur = await db.supplier_items.find_one({"product_id": si["product_id"], "supplier_id": si["supplier_id"]}, {"_id": 0})
        fields = {k: v for k, v in si.items() if v and k not in ("product_id", "supplier_id")}
        if cur:
            fields["status"] = "active"   # impor ulang sesudah rollback mengaktifkan kembali peta
            before["supplier_items"].append({"id": cur["id"], "prev": {k: cur.get(k) for k in fields}})
            await snap()
            await db.supplier_items.update_one({"id": cur["id"]}, {"$set": {**fields, "updated_at": now_iso(), "updated_by": actor}})
        else:
            nid = new_id("si")
            before["supplier_items"].append({"id": nid, "created": True})
            await snap()
            await db.supplier_items.insert_one({"id": nid, "product_id": si["product_id"], "supplier_id": si["supplier_id"],
                                                **fields, "status": "active", "created_at": now_iso(),
                                                "created_by": actor, "import_batch": batch_id})
    await db.client_dataset_imports.update_one({"id": batch_id}, {"$set": {
        "status": "committed", "committed_at": now_iso(), "before": before, "summary": fresh["summary"]}})
    return {"id": batch_id, "status": "committed", "summary": fresh["summary"]}


async def rollback(batch_id: str, actor: str) -> Dict[str, Any]:
    me = await db.client_dataset_imports.find_one({"id": batch_id}, {"_id": 0, "committed_at": 1, "status": 1})
    if not me:
        raise HTTPException(status_code=404, detail="Batch tidak ditemukan.")
    newer = None if me.get("status") == "failed" else await db.client_dataset_imports.find_one(
        {"status": "committed", "committed_at": {"$gt": me.get("committed_at") or ""}}, {"_id": 0, "filename": 1})
    if newer:
        raise HTTPException(status_code=409, detail=f"Batalkan dulu impor yang lebih baru ({newer.get('filename')}) — "
                                                    "pembatalan harus berurutan dari yang terakhir.")
    b = await db.client_dataset_imports.find_one_and_update(
        {"id": batch_id, "status": {"$in": ["committed", "failed"]}}, {"$set": {"status": "rolling_back"}})
    if not b:
        raise HTTPException(status_code=409, detail="Hanya batch yang sudah diterapkan yang bisa dibatalkan.")
    bf = b.get("before") or {}
    def _restore(prev: Dict[str, Any]) -> Dict[str, Any]:
        upd: Dict[str, Any] = {"$set": {**{k: v for k, v in prev.items() if v is not None}, "updated_at": now_iso()}}
        gone = {k: "" for k, v in prev.items() if v is None}
        if gone:
            upd["$unset"] = gone
        return upd
    for tid, prev in (bf.get("templates") or {}).items():
        await db.product_templates.update_one({"id": tid}, _restore(prev))
    for pid, prev in (bf.get("products") or {}).items():
        await db.products.update_one({"id": pid}, _restore(prev))
    for si in bf.get("supplier_items") or []:
        if si.get("created"):
            await db.supplier_items.update_one({"id": si["id"]}, {"$set": {"status": "inactive", "updated_at": now_iso()}})
        else:
            await db.supplier_items.update_one({"id": si["id"]}, {"$set": {**si["prev"], "updated_at": now_iso()}})
    await db.client_dataset_imports.update_one({"id": batch_id}, {"$set": {
        "status": "rolled_back", "rolled_back_at": now_iso(), "rolled_back_by": actor}})
    return {"id": batch_id, "status": "rolled_back"}


async def list_batches(limit: int = 30) -> List[Dict[str, Any]]:
    return await db.client_dataset_imports.find({}, {"_id": 0, "rows": 0, "plan": 0, "before": 0}).sort(
        "created_at", -1).to_list(limit)


async def get_batch(batch_id: str) -> Optional[Dict[str, Any]]:
    b = await db.client_dataset_imports.find_one({"id": batch_id}, {"_id": 0, "before": 0})
    if not b:
        raise HTTPException(status_code=404, detail="Batch tidak ditemukan.")
    b["report"] = await dry_run(b["rows"]) if b["status"] == "dry_run" else None
    if b["report"]:   # ringkasan riwayat ikut segar (mis. sesudah jenis kain ditambah)
        b["summary"], b["unknown_values"] = b["report"]["summary"], b["report"]["unknown_values"]
        await db.client_dataset_imports.update_one({"id": batch_id, "status": "dry_run"}, {"$set": {
            "summary": b["summary"], "unknown_values": b["unknown_values"]}})
    b.pop("rows", None)
    b.pop("plan", None)
    return b
