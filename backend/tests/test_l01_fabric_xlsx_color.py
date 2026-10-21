"""L01 — tambah jenis kain dari laporan impor, Excel rekap klaim, tiga warna di katalog."""
import io
import os
import uuid

import requests
from openpyxl import load_workbook
from pymongo import MongoClient

BASE = os.environ.get("REACT_APP_BACKEND_URL", "http://localhost:8001").rstrip("/")
DB = MongoClient(os.environ.get("MONGO_URL", "mongodb://localhost:27017"))[os.environ.get("DB_NAME", "test_database")]


def _h(email="admin@kainnusantara.id", pw="demo12345"):
    r = requests.post(f"{BASE}/api/auth/login", json={"email": email, "password": pw}, timeout=30)
    return {"Authorization": f"Bearer {r.json()['token']}"}


def test_add_fabric_type_and_import_accepts_it():
    adm = _h()
    code = f"tst{uuid.uuid4().hex[:5]}"
    try:
        assert requests.post(f"{BASE}/api/master/fabric-types", json={"code": code, "name": "X", "control_uom": "liter"},
                             headers=adm, timeout=30).status_code == 400
        r = requests.post(f"{BASE}/api/master/fabric-types", json={"code": code, "name": "Tes Printing", "control_uom": "meter"},
                          headers=adm, timeout=30)
        assert r.status_code == 200, r.text
        assert requests.post(f"{BASE}/api/master/fabric-types", json={"code": code, "name": "dup", "control_uom": "kg"},
                             headers=adm, timeout=30).status_code == 409
        enums = requests.get(f"{BASE}/api/enums", headers=adm, timeout=30).json()["enums"]["fabric_type"]["values"]
        assert code in [v["value"] for v in enums]
        p = DB.products.find_one({"sku": {"$regex": "^KN"}}, {"_id": 0, "sku": 1})
        csv = f"kode_artikel,nama_artikel,sku,tahap,jenis_kain\n,,{p['sku']},finished,{code}\n"
        r = requests.post(f"{BASE}/api/client-dataset/upload", files={"file": ("x.csv", csv.encode(), "text/csv")},
                          headers=adm, timeout=60)
        assert r.status_code == 200 and code not in r.json()["report"]["unknown_values"]["jenis_kain"]
        s = _h("sales.kanda@kainnusantara.id", "demo12345")
        assert requests.post(f"{BASE}/api/master/fabric-types", json={"code": code + "s", "name": "x", "control_uom": "kg"},
                             headers=s, timeout=30).status_code == 403
    finally:
        DB.md_fabric_types.delete_many({"code": {"$regex": "^tst"}})
        DB.client_dataset_imports.delete_many({"filename": "x.csv", "status": "dry_run"})


def test_recap_xlsx():
    r = requests.get(f"{BASE}/api/makloon-claims/recap.xlsx", headers=_h(), timeout=60)
    assert r.status_code == 200 and "spreadsheetml" in r.headers["content-type"]
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Ringkasan per Mitra", "Rincian"]
    s = _h("sales.kanda@kainnusantara.id", "demo12345")
    assert requests.get(f"{BASE}/api/makloon-claims/recap.xlsx", headers=s, timeout=30).status_code == 403


def test_products_carry_supplier_color_code():
    si = DB.supplier_items.find_one({"status": {"$ne": "inactive"}}, {"_id": 0, "product_id": 1})
    rows = requests.get(f"{BASE}/api/products", headers=_h(), timeout=60).json()
    rows = rows.get("items", rows) if isinstance(rows, dict) else rows
    p = next((x for x in rows if x["id"] == si["product_id"]), None)
    assert p is None or all("supplier_color_code" in c for c in p["supplier_codes"])
