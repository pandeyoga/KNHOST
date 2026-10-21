"""L11 LN-111/LN-112 + template export + cakupan penuh."""
import io
import os
import time
import uuid

import pytest
import requests
from openpyxl import load_workbook
from pymongo import MongoClient

BASE = os.environ["REACT_APP_BACKEND_URL"].rstrip("/") + "/api"
DB = MongoClient(os.environ.get("MONGO_URL", "mongodb://localhost:27017"))[os.environ.get("DB_NAME", "test_database")]
T = uuid.uuid4().hex[:6]
SO, CUST, DSG = f"so_TEST_{T}", f"cust_TEST_{T}", f"dsgn_TEST_{T}"
NUM = f"KSC/TEST_SO-{T}"


def _login(email, pw="demo12345"):
    r = requests.post(f"{BASE}/auth/login", json={"email": email, "password": pw})
    assert r.status_code == 200, r.text
    return {"Authorization": f"Bearer {r.json()['token']}", "X-Entity-Id": "ent_ksc"}


ADMIN = _login("admin@kainnusantara.id", "demo12345")


@pytest.fixture(scope="module", autouse=True)
def seed():
    items = [{"product_id": "pA", "quantity": 10, "line_total": 1_000_000, "price": 100_000},
             {"product_id": "pB", "quantity": 20, "line_total": 1_000_000, "price": 50_000}]
    DB.customers.insert_one({"id": CUST, "name": f"TEST_Cust {T}", "entity_id": "ent_ksc", "assigned_sales_id": "u_A"})
    DB.sales_orders.insert_one({"id": SO, "number": NUM, "entity_id": "ent_ksc", "customer_id": CUST,
                                "customer_name": f"TEST_Cust {T}", "sales_name": f"TEST_SalesA {T}", "status": "shipped",
                                "items": items, "net_subtotal": 2_000_000, "grand_total": 2_220_000, "order_type": "regular",
                                "created_at": "2026-09-10T03:00:00+00:00",
                                "sales_team": [{"sales_id": "u_A", "name": f"TEST_SalesA {T}", "role": "pic", "split_pct": 60},
                                               {"sales_id": "u_B", "name": f"TEST_SalesB {T}", "role": "co", "split_pct": 40}]})
    DB.shipments.insert_many([{"id": f"shp_TEST_{T}_{i}", "order_id": SO, "product_id": p, "qty": q, "status": "dispatched",
                               "entity_id": "ent_ksc"} for i, (p, q) in enumerate([("pA", 10), ("pB", 10)])])
    DB.ar_receipts.insert_many([{"id": f"arc_TEST_{T}_{i}", "number": f"TEST_AR-{T}-{i}", "entity_id": "ent_ksc", "status": "posted",
                                 "receipt_date": "2026-10-05T03:00:00+00:00", "created_at": "2026-10-05T03:00:00+00:00",
                                 "allocations": [{"order_id": SO, "order_number": NUM, "applied": a}]}
                                for i, a in enumerate([500_000, 700_000])])
    DB.sales_returns.insert_one({"id": f"sr_TEST_{T}", "order_id": SO, "total_amount": 200_000, "status": "approved",
                                 "entity_id": "ent_ksc", "created_at": "2026-10-06T03:00:00+00:00"})
    versions = [{"version": v, "by": f"TEST_DesA {T}" if v <= 8 else f"TEST_DesB {T}", "at": f"2026-10-{v:02d}T03:00:00+00:00",
                 "score": (1.5 if v % 3 == 0 else None), "score_by": "TEST_Penilai" if v % 3 == 0 else None} for v in range(1, 16)]
    DB.design_gallery.insert_one({"id": DSG, "code": f"TEST_{T}", "title": "TEST desain", "entity_id": "ent_ksc",
                                  "created_by": f"TEST_DesA {T}", "versions": versions, "version": 15,
                                  "created_at": "2026-10-01T03:00:00+00:00"})
    yield
    for c, q in (("customers", {"id": CUST}), ("sales_orders", {"id": SO}), ("shipments", {"order_id": SO}),
                 ("ar_receipts", {"id": {"$regex": f"^arc_TEST_{T}"}}), ("sales_returns", {"order_id": SO}),
                 ("design_gallery", {"id": DSG}), ("export_templates", {"name": {"$regex": "^TEST_"}})):
        DB[c].delete_many(q)


def _pv(body):
    r = requests.post(f"{BASE}/export-center/preview", json={"limit": 200, **body}, headers=ADMIN)
    assert r.status_code == 200, r.text
    return r.json()


def _so_filter():
    return {"field": "number", "cmp": "eq", "value": NUM}


def test_ln111_t01_no_multiplication():
    row = _pv({"dataset": "sales_finance", "filter": _so_filter()})["rows"]
    assert len(row) == 1
    r = row[0]
    assert r["grand_total"] == 2_220_000 and r["cash_received"] == 1_200_000 and r["returned_value"] == 200_000
    assert r["shipped_qty"] == 20 and r["shipped_value"] == 1_500_000   # 10×100rb + 10×50rb
    g = _pv({"dataset": "sales_finance", "group_by": ["customer_name"],
             "filter": {"field": "customer_name", "cmp": "eq", "value": f"TEST_Cust {T}"}})["rows"][0]
    assert g["grand_total"] == 2_220_000 and g["cash_received"] == 1_200_000 and g["_rows"] == 1


def test_ln111_t02_attribution_from_transaction():
    DB.customers.update_one({"id": CUST}, {"$set": {"assigned_sales_id": "u_Z"}})
    rows = _pv({"dataset": "sales_finance", "filter": {"op": "and", "items": [
        {"field": "sales_name", "cmp": "eq", "value": f"TEST_SalesA {T}"}, _so_filter()]}})["rows"]
    assert len(rows) == 1 and rows[0]["sales_name"] == f"TEST_SalesA {T}"


def test_ln111_t03_date_basis():
    oct_ = {"date_from": "2026-10-01", "date_to": "2026-10-31"}
    assert _pv({"dataset": "sales_finance", "filter": _so_filter(), "date_field": "created_at", **oct_})["total"] == 0
    cash = _pv({"dataset": "cash_allocations", "date_field": "receipt_date", **oct_,
                "filter": {"field": "order_number", "cmp": "eq", "value": NUM}})
    assert cash["total"] == 2 and sum(r["applied"] for r in cash["rows"]) == 1_200_000
    assert "Tanggal terima uang" in cash["date_text"]


def test_ln112_t01_t02_designer_report():
    flt = {"field": "code", "cmp": "eq", "value": f"TEST_{T}"}
    det = _pv({"dataset": "designer_report", "filter": flt})
    assert det["total"] == 15
    assert {r["author"] for r in det["rows"]} == {f"TEST_DesA {T}", f"TEST_DesB {T}"}
    assert all(r["created_by"] == f"TEST_DesA {T}" for r in det["rows"])
    summ = {r["author"]: r for r in _pv({"dataset": "designer_report", "filter": flt, "group_by": ["author"]})["rows"]}
    a, b = summ[f"TEST_DesA {T}"], summ[f"TEST_DesB {T}"]
    assert a["_rows"] == 8 and a["scored"] == 2 and a["score"] == 1.5      # v3,v6 dinilai; kosong tidak jadi nol
    assert b["_rows"] == 7 and b["scored"] == 3 and b["score"] == 1.5
    assert a["is_revision"] + b["is_revision"] == 14


def test_ln112_t03_sales_split_and_report_file():
    body = {"dataset": "sales_report", "filter": _so_filter()}
    rows = {r["person"]: r for r in _pv(body)["rows"]}
    a, b = rows[f"TEST_SalesA {T}"], rows[f"TEST_SalesB {T}"]
    assert a["attributed_net"] == 1_200_000 and b["attributed_net"] == 800_000
    assert a["attributed_net_after_returns"] == 1_080_000 and a["attributed_cash"] == 720_000
    r = requests.post(f"{BASE}/export-center/jobs", json={**body, "format": "xlsx"}, headers=ADMIN).json()
    for _ in range(60):
        j = requests.get(f"{BASE}/export-center/jobs/{r['id']}", headers=ADMIN).json()
        if j["status"] in ("done", "failed"):
            break
        time.sleep(1)
    wb = load_workbook(io.BytesIO(requests.get(f"{BASE}/export-center/jobs/{r['id']}/download", headers=ADMIN).content))
    names = wb.sheetnames
    assert names[0] == "Ringkasan" and any(n.startswith("Ringkasan per") for n in names)
    ws = wb[next(n for n in names if n.startswith("Ringkasan per"))]
    hdr = [c.value for c in ws[1]]
    tot = sum(row[hdr.index("Penjualan teratribusi")] for row in ws.iter_rows(min_row=2, values_only=True))
    assert tot == 2_000_000


def test_templates_save_run_owner_only():
    spec = {"dataset": "sales_finance", "columns": ["number", "grand_total", "cash_received"], "filter": _so_filter()}
    t = requests.post(f"{BASE}/export-center/templates", json={"name": f"TEST_tpl {T}", "spec": spec, "format": "csv"},
                      headers=ADMIN)
    assert t.status_code == 200, t.text
    tid = t.json()["id"]
    assert any(x["id"] == tid for x in requests.get(f"{BASE}/export-center/templates", headers=ADMIN).json()["items"])
    job = requests.post(f"{BASE}/export-center/templates/{tid}/run", headers=ADMIN)
    assert job.status_code == 200 and job.json()["total_rows"] == 1
    other = _login("manager@kainnusantara.id")
    assert requests.post(f"{BASE}/export-center/templates/{tid}/run", headers=other).status_code == 404
    s = _login("sales.kanda@kainnusantara.id")
    bad = requests.post(f"{BASE}/export-center/templates", headers=s,
                        json={"name": "TEST_bad", "spec": {"dataset": "inventory_rolls", "columns": ["grade"]}})
    assert bad.status_code == 403
    assert requests.delete(f"{BASE}/export-center/templates/{tid}", headers=ADMIN).status_code == 200


def test_coverage_no_unmapped():
    c = requests.get(f"{BASE}/export-center/coverage", headers=ADMIN).json()
    assert c["summary"]["unmapped"] == 0, [r["collection"] for r in c["collections"] if r["status"] == "unmapped"]
