"""L11 LN-110/LN-113 — Pusat Export: preview = file, cakupan registry, partisi, field security, unduh ber-izin ulang."""
import csv
import io
import os
import time
import uuid

import pytest
import requests
from openpyxl import load_workbook
from pymongo import MongoClient

BASE = os.environ["REACT_APP_BACKEND_URL"].rstrip("/") + "/api"
DB = MongoClient(os.environ.get("MONGO_URL", "mongodb://localhost:27017"))[os.environ.get("DB_NAME", "test_database")]
TAG = uuid.uuid4().hex[:6]
CUST = f"cust_TEST_exp_{TAG}"


def _login(email, pw="demo1234", ent="ent_ksc"):
    r = requests.post(f"{BASE}/auth/login", json={"email": email, "password": pw})
    assert r.status_code == 200, r.text
    return {"Authorization": f"Bearer {r.json()['token']}", "X-Entity-Id": ent}


ADMIN = _login("admin@kainnusantara.id", "demo12345")


@pytest.fixture(scope="module", autouse=True)
def seed():
    DB.customers.insert_one({"id": CUST, "code": f"TEST_{TAG}", "name": '=HYPERLINK("http://evil","klik")',
                             "entity_id": "ent_ksc", "status": "active", "created_at": "2026-10-01T03:00:00+00:00"})
    yield
    DB.customers.delete_many({"id": CUST})
    DB.permission_settings.update_one({"id": "default"}, {"$addToSet": {"matrix.manager.field_visibility": "grade"}})


def _job(headers, body, timeout=60):
    r = requests.post(f"{BASE}/export-center/jobs", json=body, headers=headers)
    assert r.status_code == 200, r.text
    jid = r.json()["id"]
    for _ in range(timeout):
        j = requests.get(f"{BASE}/export-center/jobs/{jid}", headers=headers).json()
        if j["status"] in ("done", "failed"):
            return j
        time.sleep(1)
    raise AssertionError("job tidak selesai")


def test_t01_preview_equals_export_with_nested_filter():
    body = {"dataset": "sales_orders", "columns": ["number", "status", "customer_name", "grand_total", "customer.code"],
            "filter": {"op": "and", "items": [{"field": "grand_total", "cmp": "gt", "value": 0},
                                              {"op": "or", "items": [{"field": "status", "cmp": "eq", "value": "done"},
                                                                     {"field": "customer_name", "cmp": "contains", "value": "toko"}]}]},
            "sort": [{"key": "grand_total", "dir": "desc"}], "date_field": "created_at",
            "date_from": "2026-01-01", "date_to": "2026-12-31", "entity_scope": "all"}
    hdr = {**ADMIN, "X-Entity-Id": "all"}
    p = requests.post(f"{BASE}/export-center/preview", json={**body, "limit": 200}, headers=hdr).json()
    assert p["total"] == len(p["rows"]) and p["total"] > 0
    assert "DAN (" in p["filter_text"]
    j = _job(hdr, {**body, "format": "xlsx"})
    assert j["status"] == "done" and j["written_rows"] == p["total"]
    f = requests.get(f"{BASE}/export-center/jobs/{j['id']}/download", headers=hdr)
    assert f.status_code == 200
    ws = load_workbook(io.BytesIO(f.content))["Data 1"]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert [r[0] for r in rows] == [x["number"] for x in p["rows"]]
    assert [r[3] for r in rows] == [x["grand_total"] for x in p["rows"]]
    assert isinstance(rows[0][3], (int, float))


def test_t02_coverage_lists_every_collection():
    c = requests.get(f"{BASE}/export-center/coverage", headers=ADMIN).json()
    names = set(DB.list_collection_names())
    assert {r["collection"] for r in c["collections"]} >= {n for n in names if not n.startswith("system.")}
    by = {r["collection"]: r for r in c["collections"]}
    assert by["sessions"]["status"] == "excluded" and by["users"]["reason"]
    assert by["sales_orders"]["status"] == "mapped"
    assert all(r.get("reason") or r.get("datasets") for r in c["collections"])


def test_t03_partitions_not_truncated():
    body = {"dataset": "inventory_movements", "columns": ["timestamp", "movement_type", "quantity"],
            "format": "xlsx", "rows_per_sheet": 5000}
    total = requests.post(f"{BASE}/export-center/preview", json={**body, "limit": 1}, headers=ADMIN).json()["total"]
    assert total > 5000
    j = _job(ADMIN, body, timeout=120)
    assert j["written_rows"] == total and len(j["partitions"]) == -(-total // 5000)
    wb = load_workbook(io.BytesIO(requests.get(f"{BASE}/export-center/jobs/{j['id']}/download", headers=ADMIN).content),
                       read_only=True)
    assert wb.sheetnames[0] == "Ringkasan"
    assert sum(sum(1 for _ in ws.iter_rows(min_row=2)) for ws in wb.worksheets[1:]) == total


def test_ln113_t01_sales_cannot_inject_hidden_fields():
    s = _login("sales.kanda@kainnusantara.id", ent="ent_kanda")
    for body in ({"dataset": "inventory_rolls", "columns": ["roll_no", "grade"]},
                 {"dataset": "inventory_rolls", "columns": ["roll_no", "lot"]},
                 {"dataset": "inventory_rolls", "columns": ["roll_no", "unit_cost"]},
                 {"dataset": "sales_order_items", "columns": ["number", "product.grade"]},
                 {"dataset": "inventory_rolls", "columns": ["roll_no"], "filter": {"field": "grade", "cmp": "eq", "value": "A"}},
                 {"dataset": "inventory_rolls", "columns": ["roll_no"], "sort": [{"key": "lot", "dir": "asc"}]}):
        for url in ("preview", "jobs"):
            r = requests.post(f"{BASE}/export-center/{url}", json=body, headers=s)
            assert r.status_code == 403 and r.json()["detail"]["code"] == "FIELD_NOT_ALLOWED", (url, body, r.text)
    r = requests.post(f"{BASE}/export-center/preview", json={"dataset": "purchase_orders"}, headers=s)
    assert r.status_code == 403
    p = requests.post(f"{BASE}/export-center/preview", json={"dataset": "sales_order_items", "columns": ["sku", "product_name"]},
                      headers=s).json()
    assert all("••" in (r["sku"] or "••") for r in p["rows"])
    ds = requests.get(f"{BASE}/export-center/datasets", headers=s).json()
    rolls = next(d for d in ds["datasets"] if d["id"] == "inventory_rolls")
    assert not {"grade", "lot", "unit_cost"} & {c["key"] for c in rolls["columns"]}


def test_ln113_t02_download_rechecked_and_owner_only():
    m = _login("manager@kainnusantara.id")
    j = _job(m, {"dataset": "inventory_rolls", "columns": ["roll_no", "grade"], "format": "csv",
                 "filter": {"field": "roll_no", "cmp": "starts", "value": "RL-0000"}})
    assert requests.get(f"{BASE}/export-center/jobs/{j['id']}/download", headers=m).status_code == 200
    other = _login("finance@kainnusantara.id")
    assert requests.get(f"{BASE}/export-center/jobs/{j['id']}/download", headers=other).status_code == 404
    DB.permission_settings.update_one({"id": "default"}, {"$pull": {"matrix.manager.field_visibility": "grade"}})
    try:
        r = requests.get(f"{BASE}/export-center/jobs/{j['id']}/download", headers=m)
        assert r.status_code == 403 and r.json()["detail"]["code"] == "ACCESS_CHANGED"
        rr = requests.post(f"{BASE}/export-center/jobs/{j['id']}/retry", headers=m)
        assert rr.status_code == 403  # retry tidak memperluas scope
    finally:
        DB.permission_settings.update_one({"id": "default"}, {"$addToSet": {"matrix.manager.field_visibility": "grade"}})


def test_ln113_t03_formula_cells_are_text():
    body = {"dataset": "customers", "columns": ["code", "name"],
            "filter": {"field": "code", "cmp": "eq", "value": f"TEST_{TAG}"}}
    jx = _job(ADMIN, {**body, "format": "xlsx"})
    wb = load_workbook(io.BytesIO(requests.get(f"{BASE}/export-center/jobs/{jx['id']}/download", headers=ADMIN).content))
    cell = wb["Data 1"]["B2"]
    assert cell.data_type == "s" and cell.value.startswith("=HYPERLINK")
    jc = _job(ADMIN, {**body, "format": "csv"})
    text = requests.get(f"{BASE}/export-center/jobs/{jc['id']}/download", headers=ADMIN).content.decode("utf-8-sig")
    row = list(csv.reader(io.StringIO(text)))[1]
    assert row[1].startswith("'=")
