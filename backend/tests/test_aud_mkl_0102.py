"""AUD-MKL-01 / AUD-MKL-02 (review 2026-10-19) — rekap susut & klaim mitra makloon.

MKL-01: Excel rekap 500 (KeyError `loss_claim_receivable`) begitu ada baris mitra → kolom kini memakai kontrak data
JSON (`loss_claim`, `loss_deducted`); angka tiap baris & total Excel dicocokkan dengan JSON dan dokumen sumber.
MKL-02: rincian dipotong diam-diam 1.000 & order dibaca maks 5.000 → ringkasan & rincian kini dari himpunan utuh;
layar memakai halaman rincian + `rows_total`, Excel memuat semua rincian.
Fixture: order makloon sintetis per mitra uji (id `mko_TEST_MKL_*`), diisolasi lewat filter `makloon_id`.
Jalankan: cd /app/backend && python -m pytest tests/test_aud_mkl_0102.py -p no:randomly -n 0 -q
"""
import io
import os
import uuid

import pytest
import requests
from openpyxl import load_workbook
from pymongo import MongoClient

BASE = os.environ.get("REACT_APP_BACKEND_URL", "http://localhost:8001").rstrip("/") + "/api"
DB = MongoClient(os.environ.get("MONGO_URL", "mongodb://localhost:27017"))[os.environ.get("DB_NAME", "test_database")]
TAG = uuid.uuid4().hex[:6]
AT = "2026-10-15T08:00:00"
BILL = f"vb_TEST_MKL_{TAG}"


def _mid(k):
    return f"mkl_TEST_{TAG}_{k}"


def _loss(status, qty, expense=0.0, claim=0.0, settle="", at=AT, bill_id=""):
    return {"status": status, "qty": qty, "over_qty": qty if status == "pending_approval" else 0,
            "expense_value": expense, "claim_value": claim, "settle": settle, "recorded_at": at,
            "decided_at": at if status == "posted" else "", "reason": f"uji {TAG}", "bill_id": bill_id}


def _order(k, losses=(), claim=None):
    n = uuid.uuid4().hex[:8]
    step = {"seq": 1, "makloon_id": _mid(k), "makloon_name": f"Mitra Uji {k} {TAG}", "input_unit": "kg",
            "output_unit": "yard", "material_losses": list(losses)}
    if claim:
        step["claim"] = claim
    return {"id": f"mko_TEST_MKL_{TAG}_{n}", "mko_number": f"MKO-T{TAG}-{n}", "entity_id": "ent_ksc",
            "status": "in_progress", "steps": [step]}


def _claim(status, amount, action="", at=AT):
    return {"status": status, "amount": amount, "action": action, "variance_qty": 3,
            "approved_at": at if status == "approved" else "", "proposed_at": at, "reason": f"selisih {TAG}"}


@pytest.fixture(scope="module")
def adm():
    r = requests.post(f"{BASE}/auth/login", json={"email": "admin@kainnusantara.id", "password": "demo12345"}, timeout=30)
    assert r.status_code == 200, r.text
    return {"Authorization": f"Bearer {r.json()['token']}"}


@pytest.fixture(scope="module", autouse=True)
def seed():
    DB.vendor_bills.insert_one({"id": BILL, "bill_number": f"BILL-T{TAG}"})
    DB.makloon_orders.insert_many([
        _order("piutang", [_loss("posted", 5, 100, 300, "tagih_ganti")]),
        _order("potong", [_loss("posted", 2, 50, 200, "potong_bon", bill_id=BILL)]),
        _order("campur", [_loss("posted", 4, 30, 100, "tagih_ganti"), _loss("posted", 1, 10, 40, "potong_bon"),
                          _loss("pending_approval", 2, at="2026-10-16T09:00:00")]),
        _order("campur", claim=_claim("approved", 150, "tagih_ganti")),
        _order("campur", claim=_claim("approved", 60, "potong_bon")),
        _order("campur", claim=_claim("pending_approval", 70)),
        _order("seribu", [_loss("posted", 1, 0, 10, "tagih_ganti", at=f"2026-10-{1 + i % 28:02d}T08:00:00")
                          for i in range(1001)]),
    ])
    DB.makloon_orders.insert_many([_order("limaribu", [_loss("posted", 1, 0, 1, "tagih_ganti")]) for _ in range(5001)])
    yield
    DB.makloon_orders.delete_many({"id": {"$regex": f"^mko_TEST_MKL_{TAG}_"}})
    DB.vendor_bills.delete_one({"id": BILL})


def _json(adm, k, **params):
    r = requests.get(f"{BASE}/makloon-claims/recap", headers=adm, params={"makloon_id": _mid(k), **params}, timeout=120)
    assert r.status_code == 200, r.text[:300]
    return r.json()


def _xlsx(adm, k):
    r = requests.get(f"{BASE}/makloon-claims/recap.xlsx", headers=adm, params={"makloon_id": _mid(k)}, timeout=180)
    assert r.status_code == 200, r.text[:300]
    return load_workbook(io.BytesIO(r.content))


def _detail_rows(wb):
    return [row for ws in wb.worksheets[1:] for row in ws.iter_rows(min_row=2, values_only=True)]


SUMMARY_KEYS = ["loss_events", "loss_qty", "loss_expense", "loss_claim", "loss_deducted", "variance_claims",
                "variance_claim_amount", "loss_pending", "variance_pending"]


def _check_summary(wb, data):
    ws = wb["Ringkasan per Mitra"]
    rows = list(ws.iter_rows(min_row=3, values_only=True))
    partners, total = rows[:len(data["partners"])], rows[len(data["partners"])]
    for got, p in zip(partners, data["partners"]):
        assert got[0] == p["makloon_name"]
        assert list(got[1:10]) == pytest.approx([p[k] for k in SUMMARY_KEYS])
        assert got[10] == pytest.approx(p["receivable_total"] + p["deducted_total"])
    t = data["totals"]
    assert total[0] == "TOTAL"
    assert list(total[1:10]) == pytest.approx([t[k] for k in SUMMARY_KEYS])
    assert total[10] == pytest.approx(t["receivable_total"] + t["deducted_total"])


# ─── AUD-MKL-01 ────────────────────────────────────────────────────────────────
@pytest.mark.parametrize("kind,receivable,deducted", [("piutang", 300, 0), ("potong", 0, 200), ("campur", 250, 100)])
def test_excel_baris_mitra_cocok_dengan_json_dan_sumber(adm, kind, receivable, deducted):
    data = _json(adm, kind)
    assert len(data["partners"]) == 1
    p = data["partners"][0]
    assert (p["receivable_total"], p["deducted_total"]) == (receivable, deducted)        # dokumen sumber (fixture)
    wb = _xlsx(adm, kind)
    _check_summary(wb, data)
    det = _detail_rows(wb)
    assert len(det) == data["rows_total"] == len(data["rows"])
    assert [(r[7], r[8], r[11]) for r in det] == [(x["expense"], x["claim"], x["status"]) for x in data["rows"]]
    if kind == "potong":
        assert det[0][10] == f"BILL-T{TAG}"                                               # nomor tagihan dipotong


def test_excel_campuran_angka_rinci(adm):
    p = _json(adm, "campur")["partners"][0]
    assert (p["loss_claim"], p["loss_deducted"], p["variance_claims"], p["variance_claim_amount"]) == (100, 40, 2, 210)
    assert (p["loss_pending"], p["variance_pending"], p["loss_events"], p["loss_expense"]) == (1, 1, 2, 40)


def test_excel_tanpa_data(adm):
    data = _json(adm, "tidak_ada")
    assert data["partners"] == [] and data["rows_total"] == 0
    wb = _xlsx(adm, "tidak_ada")
    _check_summary(wb, data)
    assert _detail_rows(wb) == []


# ─── AUD-MKL-02 ────────────────────────────────────────────────────────────────
def test_lebih_dari_seribu_kejadian_rincian_tidak_dipotong_diam_diam(adm):
    first = _json(adm, "seribu")
    assert first["totals"]["receivable_total"] == 10010 and first["rows_total"] == 1001
    assert len(first["rows"]) == 1000 and first["rows_offset"] == 0
    rest = _json(adm, "seribu", rows_offset=1000)
    assert len(rest["rows"]) == 1 and rest["rows_total"] == 1001
    assert sum(r["claim"] for r in first["rows"] + rest["rows"]) == pytest.approx(10010)
    wb = _xlsx(adm, "seribu")
    det = _detail_rows(wb)
    assert len(det) == 1001 and sum(r[8] for r in det) == pytest.approx(10010)
    _check_summary(wb, first)


def test_lebih_dari_lima_ribu_order(adm):
    data = _json(adm, "limaribu", rows_limit=10)
    assert data["totals"]["receivable_total"] == 5001 and data["totals"]["loss_events"] == 5001
    assert data["rows_total"] == 5001 and len(data["rows"]) == 10
    det = _detail_rows(_xlsx(adm, "limaribu"))
    assert len(det) == 5001 and sum(r[8] for r in det) == pytest.approx(5001)


def test_parameter_halaman_tidak_sah_ditolak(adm):
    for params in ({"rows_limit": 0}, {"rows_limit": 5000}, {"rows_offset": -1}):
        r = requests.get(f"{BASE}/makloon-claims/recap", headers=adm, params=params, timeout=60)
        assert r.status_code == 400, (params, r.status_code)
