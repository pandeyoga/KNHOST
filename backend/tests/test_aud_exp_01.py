"""AUD-EXP-01 (review 2026-10-19) — ekspor berkelompok selesai tetapi unduhan ditolak 403 ACCESS_CHANGED `_rows`.

Kolom turunan engine (`_rows` = hitungan grouping) kini dikenali asalnya; pemeriksaan ulang hak saat unduh
TETAP berlaku untuk kolom dataset & badan usaha.
Jalankan: cd /app/backend && python -m pytest tests/test_aud_exp_01.py -p no:randomly -n 0 -q
"""
import io
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import requests  # noqa: E402
from openpyxl import load_workbook  # noqa: E402
import test_l11_export_center as m  # noqa: E402

BASE, DB = m.BASE, m.DB
GROUP = {"dataset": "sales_finance", "columns": ["sales_name", "grand_total"], "group_by": ["sales_name"]}


def _dl(job, headers):
    return requests.get(f"{BASE}/export-center/jobs/{job['id']}/download", headers=headers)


def test_grouping_csv_dan_xlsx_bisa_diunduh():
    hdr = {**m.ADMIN, "X-Entity-Id": "all"}
    for fmt in ("csv", "xlsx"):
        j = m._job(hdr, {**GROUP, "format": fmt, "entity_scope": "all"})
        assert j["status"] == "done", j
        assert "_rows" in j["columns"]
        r = _dl(j, hdr)
        assert r.status_code == 200, r.text[:300]
        if fmt == "csv":
            assert "Jumlah baris" in r.content.decode("utf-8-sig").splitlines()[0]
        else:
            wb = load_workbook(io.BytesIO(r.content))
            heads = [c.value for c in next(wb.worksheets[0].iter_rows(max_row=1))]
            assert "Jumlah baris" in " ".join(str(h) for h in heads if h) or any(
                "Jumlah baris" in str(c.value) for ws in wb.worksheets for row in ws.iter_rows(max_row=12) for c in row)


def test_hak_kolom_pengelompokan_dicabut_tetap_ditolak():
    mgr = m._login("manager@kainnusantara.id")
    j = m._job(mgr, {"dataset": "inventory_rolls", "columns": ["grade", "length_remaining"], "group_by": ["grade"],
                     "format": "csv"})
    assert j["status"] == "done", j
    assert _dl(j, mgr).status_code == 200
    DB.permission_settings.update_one({"id": "default"}, {"$pull": {"matrix.manager.field_visibility": "grade"}})
    try:
        r = _dl(j, mgr)
        assert r.status_code == 403 and r.json()["detail"]["code"] == "ACCESS_CHANGED"
        assert r.json()["detail"]["fields"] == ["grade"]          # yang hilang = kolom dataset, bukan `_rows`
    finally:
        DB.permission_settings.update_one({"id": "default"}, {"$addToSet": {"matrix.manager.field_visibility": "grade"}})
