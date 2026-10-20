"""AUD-EXP-02 (review 2026-10-19) — export menggandakan baris saat sumber berubah, walau jumlah baris cocok.

Kini sumber dibaca SEKALI ke snapshot (`export_snap_<job>`, unik per dokumen sumber + baris unwind); file, ringkasan
dan nilai join dibaca dari snapshot itu. Uji memakai engine sebenarnya atas koleksi uji terisolasi; perubahan sumber
disuntikkan DI ANTARA batch penulisan (insert, hapus, ubah field filter, ubah field urut, ubah nilai, ubah join).
Pada kode lama (fetch skip/limit per batch) penyuntik yang sama dipasang pada `fetch` — itulah bukti merahnya.
Uji tingkat layanan memakai event loop BERSAMA `tests/_shared_loop.py` (klien Motor modul `db` terikat ke loop pertama).
Jalankan: cd /app/backend && python -m pytest tests/test_aud_exp_02.py -p no:randomly -n 0 -q
"""
import csv
import os
import sys
import uuid

import pytest
from openpyxl import load_workbook
from pymongo import MongoClient

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from _shared_loop import run  # noqa: E402

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from services import export_engine as eng  # noqa: E402

DB = MongoClient(os.environ.get("MONGO_URL", "mongodb://localhost:27017"))[os.environ.get("DB_NAME", "test_database")]
TAG = uuid.uuid4().hex[:6]
SRC, CUST = f"test_exp02_{TAG}", f"test_exp02c_{TAG}"
TMP = os.path.join("/tmp", f"exp02_{TAG}")


COLS = [{"key": "id", "label": "ID", "type": "text", "path": "id", "source": "base"},
        {"key": "grp", "label": "Grup", "type": "text", "path": "grp", "source": "base"},
        {"key": "rank", "label": "Urut", "type": "number", "path": "rank", "source": "base"},
        {"key": "val", "label": "Nilai", "type": "number", "path": "val", "source": "base"},
        {"key": "cust_name", "label": "Pelanggan", "type": "text", "path": "name", "source": "cust"}]
DS = {"id": "uji_exp02", "collection": SRC, "key": "id", "label": "Uji EXP-02", "grain": "1 baris = 1 dokumen",
      "joins": [{"id": "cust", "collection": CUST, "local": "cust", "foreign": "id", "fields": [{"path": "name"}]}]}


def _plan(sort=None, summary=False, cols=COLS, ds=DS, pipeline=None):
    p = {"ds": ds, "columns": cols, "group": None, "pipeline": pipeline or [{"$match": {"grp": "A"}}],
         "sort": sort or {"id": 1}, "hidden": None, "filter_text": "grp = A", "date_text": "(semua tanggal)",
         "sort_text": "uji", "entity_ids": ["ent_ksc"], "denied": []}
    if summary:
        base = {c["key"]: c for c in COLS if c["source"] == "base"}
        scols, group = eng._group_plan(DS, ["grp"], [base["grp"], base["val"]], base)
        p["summary_plan"] = {"ds": DS, "columns": scols, "group": group, "pipeline": p["pipeline"], "sort": {},
                             "hidden": None}
    return p


def _job(fmt="csv"):
    return {"id": f"exj_TEST_{TAG}_{uuid.uuid4().hex[:6]}", "format": fmt, "created_at": "2026-10-21T00:00:00",
            "user_name": "uji", "role": "admin", "written_rows": 0}


@pytest.fixture(autouse=True)
def fresh():
    DB[SRC].drop()
    DB[CUST].drop()
    DB[SRC].insert_many([{"id": str(i), "grp": "A", "rank": i, "val": 10.0 * i, "cust": "c1"} for i in range(1, 5)]
                        + [{"id": "5", "grp": "B", "rank": 5, "val": 50.0, "cust": "c1"}])
    DB[CUST].insert_one({"id": "c1", "name": "Pelanggan Lama"})
    os.makedirs(TMP, exist_ok=True)
    yield
    DB[SRC].drop()
    DB[CUST].drop()


def _inject(monkeypatch, mutate):
    """Jalankan `mutate` sesudah batch PERTAMA ditulis (kode baru: cursor snapshot; kode lama: fetch per halaman)."""
    orig_batches = getattr(eng, "_snap_batches", None)
    if orig_batches:
        async def batches(*a, **kw):
            n = 0
            async for chunk in orig_batches(*a, **kw):
                yield chunk
                n += 1
                if n == 1:
                    mutate()
        monkeypatch.setattr(eng, "_snap_batches", batches)
    orig_fetch, st = eng.fetch, {"n": 0}

    async def fetch(plan, skip, limit):
        rows = await orig_fetch(plan, skip, limit)
        st["n"] += 1
        if st["n"] == 1:
            mutate()
        return rows
    monkeypatch.setattr(eng, "fetch", fetch)


def _csv(plan, monkeypatch=None, mutate=None):
    if mutate:
        _inject(monkeypatch, mutate)
    job = _job()
    path = os.path.join(TMP, f"{job['id']}.csv")
    res = run(eng.write_file(job, plan, path, 1000, batch=2))
    with open(path, encoding="utf-8-sig") as fh:
        rows = list(csv.DictReader(fh))
    return res, rows, job


@pytest.mark.parametrize("case", ["insert_hapus", "ubah_filter", "ubah_nilai"])
def test_perubahan_sumber_saat_export_tidak_menggandakan_atau_menggeser(monkeypatch, case):
    def mutate():
        if case == "insert_hapus":                      # skenario reviewer: hapus 4, tambah 0
            DB[SRC].delete_one({"id": "4"})
            DB[SRC].insert_one({"id": "0", "grp": "A", "rank": 0, "val": 0.0, "cust": "c1"})
        elif case == "ubah_filter":
            DB[SRC].update_one({"id": "3"}, {"$set": {"grp": "B"}})
            DB[SRC].update_one({"id": "5"}, {"$set": {"grp": "A"}})
        else:
            DB[SRC].update_one({"id": "4"}, {"$set": {"val": 999.0}})
            DB[SRC].update_one({"id": "1"}, {"$set": {"val": 999.0}})
    res, rows, _ = _csv(_plan(), monkeypatch, mutate)
    assert [r["ID"] for r in rows] == ["1", "2", "3", "4"], rows
    assert res["written_rows"] == 4
    assert [float(r["Nilai"]) for r in rows] == [10, 20, 30, 40]      # nilai dari satu versi (snapshot)


def test_perubahan_field_urut_saat_export(monkeypatch):
    res, rows, _ = _csv(_plan(sort={"rank": 1, "id": 1}), monkeypatch,
                        lambda: DB[SRC].update_many({"id": {"$in": ["1", "2"]}}, {"$inc": {"rank": 100}}))
    assert [r["ID"] for r in rows] == ["1", "2", "3", "4"] and res["written_rows"] == 4


def test_nilai_join_satu_versi_untuk_seluruh_file(monkeypatch):
    _, rows, _ = _csv(_plan(), monkeypatch, lambda: DB[CUST].update_one({"id": "c1"}, {"$set": {"name": "Pelanggan Baru"}}))
    assert {r["Pelanggan"] for r in rows} == {"Pelanggan Lama"}, rows


def test_ringkasan_dari_snapshot_yang_sama_dengan_rincian(monkeypatch):
    _inject(monkeypatch, lambda: DB[SRC].update_many({}, {"$set": {"val": 1.0}}))
    job = _job("xlsx")
    path = os.path.join(TMP, f"{job['id']}.xlsx")
    res = run(eng.write_file(job, _plan(summary=True), path, 1000, batch=2))
    wb = load_workbook(path)
    data = list(wb["Data 1"].iter_rows(min_row=2, values_only=True))
    detail_sum = sum(r[3] for r in data)
    summary = [ws for ws in wb.worksheets if ws.title.startswith("Ringkasan per")][0]
    srow = list(summary.iter_rows(min_row=2, values_only=True))
    assert len(data) == 4 and detail_sum == 100
    assert len(srow) == 1 and srow[0][0] == "A" and srow[0][1] == pytest.approx(detail_sum) and srow[0][2] == 4
    meta = {r[0]: r[1] for r in wb["Ringkasan"].iter_rows(values_only=True) if r and r[0]}
    assert meta["Total baris"] == 4 and str(meta.get("Data per", "")).startswith(res["snapshot_at"])


def test_baris_unwind_tetap_terpisah():
    DB[SRC].update_one({"id": "1"}, {"$set": {"items": [{"sku": "X"}, {"sku": "Y"}]}})
    ds = {**DS, "unwind": "items"}
    cols = [COLS[0], {"key": "sku", "label": "SKU", "type": "text", "path": "items.sku", "source": "base"}]
    pipe = [{"$match": {"id": "1"}}, {"$unwind": {"path": "$items", "includeArrayIndex": "_line"}}]
    _, rows, _ = _csv(_plan(sort={"id": 1, "_line": 1}, cols=cols, ds=ds, pipeline=pipe))
    assert [(r["ID"], r["SKU"]) for r in rows] == [("1", "X"), ("1", "Y")]


def test_snapshot_dibuang_sesudah_selesai_dan_saat_gagal(monkeypatch):
    _, _, job = _csv(_plan())
    assert f"export_snap_{job['id']}" not in DB.list_collection_names()

    async def boom(*a, **kw):
        raise RuntimeError("gangguan tulis (uji)")
    monkeypatch.setattr(eng, "_shape", boom)
    job2 = _job()
    with pytest.raises(RuntimeError):
        run(eng.write_file(job2, _plan(), os.path.join(TMP, "x.csv"), 1000, batch=2))
    assert f"export_snap_{job2['id']}" not in DB.list_collection_names()
